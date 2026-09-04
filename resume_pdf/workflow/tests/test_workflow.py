from __future__ import annotations

import io
import json
from pathlib import Path

import pytest
from docx import Document

from resume_pdf.workflow.renderer import UNACHIEVED_LABEL, render_docx, render_html
from resume_pdf.workflow.validation import ResumeWorkflowError, validate_content

EXAMPLE = Path(__file__).parents[1] / "examples" / "fictional-target.json"


def example_content() -> dict:
    return json.loads(EXAMPLE.read_text(encoding="utf-8"))


def test_example_target_is_valid_and_uses_traceable_learning_marker():
    content, summary = validate_content(example_content())

    assert content["kind"] == "target"
    assert summary == {
        "sections": 5,
        "bullets": 7,
        "unachieved": 1,
        "estimated_lines": summary["estimated_lines"],
    }
    assert summary["estimated_lines"] <= 48


def test_submittable_rejects_unachieved_target():
    content = example_content()
    content["kind"] = "submittable"

    with pytest.raises(ResumeWorkflowError) as captured:
        validate_content(content)

    assert any("submittable bullets must be true" in issue for issue in captured.value.issues)


def test_achieved_bullet_requires_fact_source_and_no_new_claims():
    content = example_content()
    bullet = content["sections"][1]["items"][0]["bullets"][0]
    bullet["source_fact_ids"] = []
    bullet["new_claims"] = ["unsupported"]

    with pytest.raises(ResumeWorkflowError) as captured:
        validate_content(content)

    assert any("source_fact_ids: required" in issue for issue in captured.value.issues)
    assert any("new_claims: must be empty" in issue for issue in captured.value.issues)


def test_docx_uses_a4_layout_and_visible_unachieved_marker():
    content, _summary = validate_content(example_content())
    document = Document(io.BytesIO(render_docx(content)))
    section = document.sections[0]
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    text += "\n" + "\n".join(
        paragraph.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
        for paragraph in cell.paragraphs
    )

    assert round(section.page_width.mm) == 210
    assert round(section.page_height.mm) == 297
    assert UNACHIEVED_LABEL in text
    assert "求职方向：数据运营助理" in text


def test_html_is_print_ready_and_escapes_resume_text():
    content, _summary = validate_content(example_content())
    content["sections"][1]["items"][0]["bullets"][0]["text"] = "整理 <script>"

    html = render_html(content)

    assert "@page{size:A4;margin:0}" in html
    assert UNACHIEVED_LABEL in html
    assert "整理 &lt;script&gt;" in html
    assert "整理 <script>" not in html


def test_education_is_always_rendered_directly_after_headline():
    content, _summary = validate_content(example_content())
    html = render_html(content)
    document = Document(io.BytesIO(render_docx(content)))
    body_xml = document._element.body.xml

    assert html.index("教育背景") < html.index("对照 JD 尚未具备")
    assert body_xml.index("教育背景") < body_xml.index("对照 JD 尚未具备")
