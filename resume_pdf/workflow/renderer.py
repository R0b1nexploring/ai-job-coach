from __future__ import annotations

import base64
import io
import json
import mimetypes
import re
import shutil
import subprocess
from html import escape
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

PROFILE_PATH = Path(__file__).with_name("template_profile.json")
UNACHIEVED_LABEL = "[未达成 · 学习目标]"


def load_profile() -> dict:
    return json.loads(PROFILE_PATH.read_text(encoding="utf-8"))


def render_docx(content: dict, photo_path: Path | None = None) -> bytes:
    profile = load_profile()
    page = profile["page"]
    typography = dict(profile["typography"])
    typography["cjk_font"] = _resolve_cjk_font(typography)
    colors = profile["colors"]
    doc = Document()
    doc.core_properties.title = "Tailored Resume"
    doc.core_properties.author = ""
    section = doc.sections[0]
    section.page_width = Mm(page["width_mm"])
    section.page_height = Mm(page["height_mm"])
    section.left_margin = Mm(page["margin_horizontal_mm"])
    section.right_margin = Mm(page["margin_horizontal_mm"])
    section.top_margin = Mm(page["margin_vertical_mm"])
    section.bottom_margin = Mm(page["margin_vertical_mm"])
    section.header_distance = Mm(5)
    section.footer_distance = Mm(5)

    palette = {
        key: RGBColor.from_string(value)
        for key, value in colors.items()
        if key != "target_background"
    }
    _configure_styles(doc, typography, palette)
    _render_header(doc, content, photo_path, typography, palette)
    for resume_section in _ordered_sections(content.get("sections") or []):
        _render_section(doc, resume_section, typography, palette)

    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


def render_html(content: dict, photo_path: Path | None = None) -> str:
    profile = load_profile()
    typography = dict(profile["typography"])
    typography["cjk_font"] = _resolve_cjk_font(typography)
    colors = profile["colors"]
    contact = content.get("contact") or {}
    details = []
    if contact.get("phone"):
        details.append(f"电话：{contact['phone']}")
    if contact.get("email"):
        details.append(f"邮箱：{contact['email']}")
    if contact.get("city"):
        details.append(f"城市：{contact['city']}")
    contact_html = "".join(f"<span>{escape(value)}</span>" for value in details)
    photo_html = ""
    header_class = "header with-photo" if photo_path and photo_path.is_file() else "header"
    if photo_path and photo_path.is_file():
        mime = mimetypes.guess_type(photo_path.name)[0] or "image/jpeg"
        encoded = base64.b64encode(photo_path.read_bytes()).decode("ascii")
        photo_html = f'<img class="portrait" src="data:{mime};base64,{encoded}" alt="">'
    sections_html = "".join(
        _section_html(section)
        for section in _ordered_sections(content.get("sections") or [])
        if section.get("items")
    )
    return f"""<!doctype html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<title>{escape(str(content.get("headline") or "简历"))}</title>
<style>
*{{box-sizing:border-box}}
html,body{{margin:0;padding:0;background:#fff;color:#{colors['ink']}}}
body{{font-family:"{escape(typography['cjk_font'])}","{escape(typography['latin_font'])}",sans-serif;font-size:{typography['body_size_pt']}pt;line-height:{typography['line_spacing']};-webkit-print-color-adjust:exact;print-color-adjust:exact}}
.page{{width:210mm;height:297mm;padding:12.5mm 14.5mm;overflow:hidden;background:#fff}}
.header{{display:grid;grid-template-columns:45mm 1fr;align-items:center;min-height:25mm}}
.header.with-photo{{grid-template-columns:45mm 1fr 28mm}}
.name{{margin:0;color:#{colors['accent']};font-size:{typography['name_size_pt']}pt;line-height:1.1;font-weight:700;letter-spacing:.08em}}
.contact{{display:flex;flex-direction:column;gap:2mm;font-size:{typography['meta_size_pt']}pt}}
.portrait{{display:block;width:27mm;max-height:36mm;object-fit:cover;margin-left:auto}}
.headline{{margin:1mm 0 2mm;padding:0 0 1.5mm;border-bottom:1.25pt solid #{colors['accent']};color:#{colors['accent']};font-size:11pt;font-weight:700}}
section{{margin:0 0 1.2mm;break-inside:avoid}}
h2{{margin:2.2mm 0 1mm;padding:0 0 .8mm;border-bottom:1.5pt solid #{colors['accent']};color:#{colors['accent']};font-size:{typography['section_size_pt']}pt;line-height:1.1}}
.entry{{margin:0 0 1mm;break-inside:avoid}}
.entry-head{{display:grid;grid-template-columns:40mm 1fr 56mm;align-items:start;font-weight:700;margin:0 0 .6mm}}
.entry-head span:nth-child(2){{text-align:center}}
.entry-head span:nth-child(3){{text-align:right}}
ul{{list-style:none;margin:0;padding:0}}
li{{position:relative;margin:0 0 .55mm;padding-left:5mm}}
li::before{{content:"◆";position:absolute;left:.7mm;top:.08em;color:#{colors['accent']};font-size:7.5pt}}
li.unachieved{{color:#{colors['muted']}}}
.target-label{{font-weight:700}}
.skills{{margin:0;line-height:1.35}}
.skills .unachieved{{color:#{colors['muted']}}}
@page{{size:A4;margin:0}}
</style>
</head>
<body>
<main class="page">
  <header class="{header_class}">
    <h1 class="name">{escape(_spaced_cjk_name(str(contact.get("name") or "")))}</h1>
    <div class="contact">{contact_html}</div>
    {photo_html}
  </header>
  <p class="headline">求职方向：{escape(str(content.get("headline") or ""))}</p>
  {sections_html}
</main>
</body>
</html>"""


def _ordered_sections(sections: list[dict]) -> list[dict]:
    """Keep source order except for the education block required below the headline."""
    education_names = {"教育背景", "教育经历", "学历背景"}
    education = [section for section in sections if section.get("name") in education_names]
    remaining = [section for section in sections if section.get("name") not in education_names]
    return [*education, *remaining]


def _section_html(section: dict) -> str:
    name = escape(str(section.get("name") or "相关经历"))
    items = section.get("items") or []
    if section.get("presentation") == "tags":
        tags = []
        for item in items:
            for bullet in item.get("bullets") or []:
                value = escape(str(bullet.get("text") or ""))
                if bullet.get("achieved") is False:
                    value = f'<span class="unachieved">{escape(UNACHIEVED_LABEL)} {value}</span>'
                tags.append(value)
        return f'<section><h2>{name}</h2><p class="skills">{" / ".join(tags)}</p></section>'
    entries = "".join(_item_html(item) for item in items)
    return f"<section><h2>{name}</h2>{entries}</section>"


def _item_html(item: dict) -> str:
    start = _format_date(item.get("start_date"))
    end = _format_date(item.get("end_date"))
    dates = " - ".join(part for part in (start, end) if part)
    organization = str(item.get("organization") or "").strip()
    location = str(item.get("location") or "").strip()
    if location:
        organization = f"{organization}｜{location}" if organization else location
    role = str(item.get("role_title") or "").strip()
    header = ""
    if any((dates, organization, role)):
        header = (
            '<div class="entry-head">'
            f"<span>{escape(dates)}</span><span>{escape(organization)}</span>"
            f"<span>{escape(role)}</span></div>"
        )
    bullets = []
    for bullet in item.get("bullets") or []:
        value = escape(str(bullet.get("text") or ""))
        if bullet.get("achieved") is False:
            value = f'<span class="target-label">{escape(UNACHIEVED_LABEL)}</span> {value}'
            bullets.append(f'<li class="unachieved">{value}</li>')
        else:
            bullets.append(f"<li>{value}</li>")
    return f'<div class="entry">{header}<ul>{"".join(bullets)}</ul></div>'


def _configure_styles(doc: Document, typography: dict, palette: dict) -> None:
    normal = doc.styles["Normal"]
    _style_font(
        normal,
        typography,
        size=typography["body_size_pt"],
        color=palette["ink"],
    )
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(1.5)
    normal.paragraph_format.line_spacing = typography["line_spacing"]

    styles = (
        ("Resume Name", typography["name_size_pt"], palette["accent"], True),
        ("Resume Headline", 11, palette["accent"], True),
        ("Resume Contact", typography["meta_size_pt"], palette["ink"], False),
        ("Resume Section", typography["section_size_pt"], palette["accent"], True),
        ("Resume Entry", typography["body_size_pt"], palette["ink"], True),
        ("Resume Bullet", typography["body_size_pt"], palette["ink"], False),
        ("Resume Skills", typography["body_size_pt"], palette["ink"], False),
    )
    for name, size, color, bold in styles:
        style = doc.styles.add_style(name, 1)
        _style_font(style, typography, size=size, color=color, bold=bold)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(1.5)
        style.paragraph_format.line_spacing = typography["line_spacing"]
    doc.styles["Resume Section"].paragraph_format.space_before = Pt(5)
    doc.styles["Resume Section"].paragraph_format.space_after = Pt(2)
    doc.styles["Resume Section"].paragraph_format.keep_with_next = True
    doc.styles["Resume Entry"].paragraph_format.keep_with_next = True


def _style_font(style, typography: dict, *, size: float, color, bold: bool = False) -> None:
    style.font.name = typography["latin_font"]
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = color
    fonts = style._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), typography["latin_font"])
    fonts.set(qn("w:hAnsi"), typography["latin_font"])
    fonts.set(qn("w:eastAsia"), typography["cjk_font"])


def _render_header(doc, content, photo_path, typography, palette) -> None:
    has_photo = bool(photo_path and photo_path.is_file())
    table = doc.add_table(rows=1, cols=3 if has_photo else 2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _remove_table_borders(table)
    widths = (Mm(42), Mm(110), Mm(27)) if has_photo else (Mm(45), Mm(134))
    for cell, width in zip(table.rows[0].cells, widths):
        cell.width = width
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        _set_cell_margins(cell, top=0, start=0, bottom=0, end=2)

    contact = content.get("contact") or {}
    name = _spaced_cjk_name(str(contact.get("name") or ""))
    name_paragraph = table.cell(0, 0).paragraphs[0]
    name_paragraph.style = doc.styles["Resume Name"]
    name_paragraph.add_run(name)

    details = []
    if contact.get("phone"):
        details.append(f"电话：{contact['phone']}")
    if contact.get("email"):
        details.append(f"邮箱：{contact['email']}")
    if contact.get("city"):
        details.append(f"城市：{contact['city']}")
    detail_cell = table.cell(0, 1)
    for index, text in enumerate(details):
        paragraph = detail_cell.paragraphs[0] if index == 0 else detail_cell.add_paragraph()
        paragraph.style = doc.styles["Resume Contact"]
        paragraph.add_run(text)

    if has_photo:
        paragraph = table.cell(0, 2).paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        paragraph.add_run().add_picture(str(photo_path), width=Mm(27))

    headline = str(content.get("headline") or "").strip()
    if headline:
        paragraph = doc.add_paragraph(style="Resume Headline")
        paragraph.add_run(f"求职方向：{headline}")
        _set_bottom_border(paragraph, palette["accent"], size="10")


def _render_section(doc, section: dict, typography: dict, palette: dict) -> None:
    items = section.get("items") or []
    if not items:
        return
    heading = doc.add_paragraph(style="Resume Section")
    heading.add_run(str(section.get("name") or "相关经历"))
    _set_bottom_border(heading, palette["accent"], size="12")
    if section.get("presentation") == "tags":
        _render_skills(doc, items, typography, palette)
        return
    education = str(section.get("name") or "") == "教育背景"
    for item in items:
        _render_item(doc, item, education, typography, palette)


def _render_item(doc, item: dict, education: bool, typography: dict, palette: dict) -> None:
    start = _format_date(item.get("start_date"))
    end = _format_date(item.get("end_date"))
    dates = " - ".join(part for part in (start, end) if part)
    organization = str(item.get("organization") or "").strip()
    location = str(item.get("location") or "").strip()
    if location:
        organization = f"{organization}｜{location}" if organization else location
    role = str(item.get("role_title") or "").strip()
    if any((dates, organization, role)):
        table = doc.add_table(rows=1, cols=3)
        table.autofit = False
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        _remove_table_borders(table)
        widths = (Mm(40), Mm(83), Mm(56))
        values = (dates, organization, role)
        alignments = (
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.RIGHT,
        )
        if education:
            values = (dates, organization, role)
        for cell, width, value, alignment in zip(
            table.rows[0].cells, widths, values, alignments
        ):
            cell.width = width
            _set_cell_margins(cell, top=0, start=0, bottom=0, end=0)
            paragraph = cell.paragraphs[0]
            paragraph.style = doc.styles["Resume Entry"]
            paragraph.alignment = alignment
            paragraph.add_run(value)
    for bullet in item.get("bullets") or []:
        _add_bullet(doc, bullet, typography, palette)


def _add_bullet(doc, bullet: dict, typography: dict, palette: dict) -> None:
    paragraph = doc.add_paragraph(style="Resume Bullet")
    paragraph.paragraph_format.left_indent = Mm(5)
    paragraph.paragraph_format.first_line_indent = Mm(-4)
    paragraph.paragraph_format.keep_together = True
    marker = paragraph.add_run("◆ ")
    _run_font(marker, typography, size=7.5, color=palette["accent"], bold=True)
    text = str(bullet.get("text") or "")
    if bullet.get("achieved") is False:
        label = paragraph.add_run(f"{UNACHIEVED_LABEL} ")
        _run_font(label, typography, size=typography["body_size_pt"], color=palette["muted"], bold=True)
        value = paragraph.add_run(text)
        _run_font(value, typography, size=typography["body_size_pt"], color=palette["muted"])
    else:
        value = paragraph.add_run(text)
        _run_font(value, typography, size=typography["body_size_pt"], color=palette["ink"])


def _render_skills(doc, items: list, typography: dict, palette: dict) -> None:
    paragraph = doc.add_paragraph(style="Resume Skills")
    index = 0
    for item in items:
        for bullet in item.get("bullets") or []:
            if index:
                paragraph.add_run(" / ")
            text = str(bullet.get("text") or "")
            if bullet.get("achieved") is False:
                text = f"{UNACHIEVED_LABEL} {text}"
                run = paragraph.add_run(text)
                _run_font(run, typography, size=typography["body_size_pt"], color=palette["muted"])
            else:
                run = paragraph.add_run(text)
                _run_font(run, typography, size=typography["body_size_pt"], color=palette["ink"])
            index += 1


def _run_font(run, typography: dict, *, size: float, color, bold: bool = False) -> None:
    run.font.name = typography["latin_font"]
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run._element.get_or_add_rPr().get_or_add_rFonts().set(
        qn("w:eastAsia"), typography["cjk_font"]
    )


def _format_date(value) -> str:
    text = str(value or "").strip()
    matched = re.fullmatch(r"(\d{4})-(\d{2})(?:-\d{2})?", text)
    return f"{matched.group(1)}.{matched.group(2)}" if matched else text


def _resolve_cjk_font(typography: dict) -> str:
    candidates = typography.get("cjk_font_candidates") or [typography["cjk_font"]]
    if not shutil.which("fc-match"):
        return str(candidates[0])
    for candidate in candidates:
        result = subprocess.run(
            ["fc-match", "-f", "%{family}", str(candidate)],
            capture_output=True,
            text=True,
            check=False,
        )
        families = [value.strip().casefold() for value in result.stdout.split(",")]
        if str(candidate).casefold() in families:
            return str(candidate)
    raise RuntimeError("No supported CJK font is installed")


def _spaced_cjk_name(value: str) -> str:
    compact = re.sub(r"\s+", "", value.strip())
    return " ".join(compact) if re.fullmatch(r"[\u4e00-\u9fff]{2,4}", compact) else value.strip()


def _set_bottom_border(paragraph, color, *, size: str) -> None:
    properties = paragraph._p.get_or_add_pPr()
    borders = properties.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        properties.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), str(color))
    borders.append(bottom)


def _remove_table_borders(table) -> None:
    properties = table._tbl.tblPr
    borders = properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "nil")
        borders.append(element)


def _set_cell_margins(cell, *, top: int, start: int, bottom: int, end: int) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{name}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
