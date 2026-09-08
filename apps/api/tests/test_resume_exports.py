import io
import zipfile
from xml.etree import ElementTree

from docx import Document

from app.services.docx_resume_renderer import render_resume_docx
from app.services.html_resume_renderer import render_resume_html
from app.services.resume_template import template_snapshot

WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def v2_content():
    return {
        "content_version": 2,
        "kind": "target",
        "headline": "机械设计工程师",
        "contact": {
            "name": "张三",
            "phone": "13800000000",
            "email": "zhangsan@example.test",
            "city": "杭州",
        },
        "sections": [
            {
                "name": "项目经历",
                "items": [
                    {
                        "organization": "示例大学",
                        "role_title": "项目负责人",
                        "location": "杭州",
                        "start_date": "2024-03",
                        "end_date": "2024-06",
                        "bullets": [
                            {
                                "text": "完成 CAD 机构设计",
                                "source_fact_ids": ["fact-1"],
                                "source_target_ids": [],
                                "achieved": True,
                            },
                            {
                                "text": "围绕 SolidWorks 完成装配体练习",
                                "source_fact_ids": [],
                                "source_target_ids": ["target-1"],
                                "achieved": False,
                            },
                        ],
                    }
                ],
            },
            {
                "name": "教育背景",
                "items": [
                    {
                        "organization": "示例大学",
                        "role_title": "机械工程本科",
                        "start_date": "2021-09",
                        "end_date": "2025-06",
                        "bullets": [
                            {
                                "text": "具备机械工程本科背景",
                                "source_fact_ids": ["fact-2"],
                                "source_target_ids": [],
                                "achieved": True,
                            }
                        ],
                    }
                ],
            },
        ],
        "excluded_gaps": [],
    }


def test_html_renderer_escapes_v1_resume_content_and_is_print_ready():
    content = {
        "headline": "设备工程师 <script>alert(1)</script>",
        "sections": [
            {"name": "项目经历", "items": [{"text": "完成 CAD 设计 <img src=x onerror=alert(1)>"}]}
        ],
    }
    result = render_resume_html(content, "备用标题")
    assert "<script>alert(1)</script>" not in result
    assert "&lt;script&gt;alert(1)&lt;/script&gt;" in result
    assert "&lt;img src=x onerror=alert(1)&gt;" in result
    assert "@page" in result and "size:A4" in result
    assert "scale(" not in result


def test_docx_uses_a4_and_portable_chinese_fonts():
    content = render_resume_docx(v2_content(), "备用标题")
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        document_xml = archive.read("word/document.xml")
        styles_xml = archive.read("word/styles.xml")
    root = ElementTree.fromstring(document_xml)
    page_size = root.find(f".//{{{WORD_NS}}}pgSz")
    assert page_size is not None
    assert page_size.attrib[f"{{{WORD_NS}}}w"] == "11906"
    assert page_size.attrib[f"{{{WORD_NS}}}h"] == "16838"
    styles_text = styles_xml.decode("utf-8")
    assert "微软雅黑" in styles_text
    assert "黑体" in styles_text
    assert "Calibri" in styles_text
    assert "Arial Unicode MS" not in styles_text
    assert 'w:styleId="ResumeBullet"' in styles_text
    assert 'w:basedOn w:val="ListBullet"' in styles_text
    assert b'<w:tab w:val="right"' in styles_xml


def test_skill_tags_render_inline_in_html_and_docx():
    content = {
        "content_version": 2,
        "kind": "submittable",
        "headline": "机械设计工程师",
        "contact": {"name": "张三", "phone": "13800000000", "email": "", "city": ""},
        "sections": [
            {
                "name": "技能",
                "presentation": "tags",
                "items": [
                    {
                        "bullets": [
                            {"text": "SolidWorks", "achieved": True},
                            {"text": "AutoCAD", "achieved": True},
                            {"text": "GD&T", "achieved": True},
                        ]
                    }
                ],
            }
        ],
        "excluded_gaps": [],
    }

    html = render_resume_html(content, "备用标题")
    document = Document(io.BytesIO(render_resume_docx(content, "备用标题")))
    assert "SolidWorks / AutoCAD / GD&amp;T" in html or "SolidWorks / AutoCAD / GD&T" in html
    assert 'class="skill-tags"' in html
    assert html.count("<li") == 0
    assert any("SolidWorks / AutoCAD / GD&T" in paragraph.text for paragraph in document.paragraphs)
    assert sum(paragraph.style.name == "Resume Bullet" for paragraph in document.paragraphs) == 0
    assert sum(paragraph.style.name == "Resume Skills" for paragraph in document.paragraphs) == 1


def test_v2_html_and_docx_show_same_sections_bullets_and_unachieved_marker():
    content = v2_content()
    html = render_resume_html(content, "备用标题")
    docx_bytes = render_resume_docx(content, "备用标题")
    document = Document(io.BytesIO(docx_bytes))
    section_count = sum(paragraph.style.name == "Resume Section" for paragraph in document.paragraphs)
    bullet_count = sum(paragraph.style.name == "Resume Bullet" for paragraph in document.paragraphs)
    assert html.count('data-resume-section="true"') == section_count == 2
    assert html.count("<li") == bullet_count == 3
    assert "未达成 · 学习目标" in html
    assert any("未达成 · 学习目标" in paragraph.text for paragraph in document.paragraphs)
    assert "电话：13800000000" in html
    assert "邮箱：zhangsan@example.test" in html
    assert "城市：杭州" in html
    assert "意向：机械设计工程师" in html
    assert "项目负责人｜示例大学｜2024.03–2024.06" in html
    assert any("电话：13800000000" in p.text and "邮箱：zhangsan@example.test" in p.text for p in document.paragraphs)
    assert any("意向：机械设计工程师" in p.text for p in document.paragraphs)
    assert any("项目负责人｜示例大学｜2024.03–2024.06" in p.text for p in document.paragraphs)


def test_renderers_share_header_section_order_and_skip_empty_bullets():
    content = {
        "content_version": 3,
        "kind": "submittable",
        "headline": "机械设计工程师",
        "contact": {
            "name": "张三",
            "phone": "13800000000",
            "email": "zhangsan@example.test",
            "city": "杭州",
        },
        "sections": [
            {
                "name": "项目/科研/竞赛",
                "items": [
                    {
                        "organization": "示例大学",
                        "role_title": "项目负责人",
                        "start_date": "2024-03",
                        "end_date": "2024-06",
                        "bullets": [
                            {"text": "完成 CAD 机构设计", "achieved": True},
                            {"text": "   ", "achieved": True},
                        ],
                    }
                ],
            },
            {
                "name": "教育背景",
                "items": [
                    {
                        "organization": "示例大学",
                        "role_title": "机械工程本科",
                        "start_date": "2021-09",
                        "end_date": "2025-06",
                        "bullets": [{"text": "具备机械工程本科背景", "achieved": True}],
                    }
                ],
            },
            {"name": "奖项证书", "items": []},
        ],
    }
    snapshot = template_snapshot()
    html = render_resume_html(content, "备用标题", snapshot=snapshot)
    document = Document(io.BytesIO(render_resume_docx(content, "备用标题", snapshot=snapshot)))
    texts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]

    assert html.find("<h1>") < html.find("意向：机械设计工程师") < html.find("电话：")
    assert html.find("教育背景") < html.find("项目/科研/竞赛")
    assert "奖项证书" not in html
    assert html.count("<li") == 2
    assert "完成 CAD 机构设计" in html

    name_i = next(index for index, text in enumerate(texts) if "张" in text)
    subtitle_i = next(index for index, text in enumerate(texts) if text.startswith("意向："))
    contact_i = next(index for index, text in enumerate(texts) if "电话：" in text)
    education_i = next(index for index, text in enumerate(texts) if text == "教育背景")
    project_i = next(index for index, text in enumerate(texts) if text == "项目/科研/竞赛")
    assert name_i < subtitle_i < contact_i < education_i < project_i
    assert "奖项证书" not in texts
    assert sum(paragraph.style.name == "Resume Bullet" for paragraph in document.paragraphs) == 2


def test_snapshot_keeps_heading_heiti():
    content = render_resume_docx(v2_content(), "备用标题", snapshot=template_snapshot())
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        styles_text = archive.read("word/styles.xml").decode("utf-8")
    assert "黑体" in styles_text
    assert "微软雅黑" in styles_text
