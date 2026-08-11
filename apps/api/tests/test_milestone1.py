import io
from datetime import datetime, timedelta, timezone

from conftest import docx_file, pdf_file, register
from docx import Document

from app.models.entities import Job, UserProfile
from app.services.job_normalizer import normalize_job
from app.services.matcher import score_match
from app.services.resume_parser import sanitize_extraction


def test_resume_upload_type_and_user_isolation(client):
    register(client)
    uploaded = client.post("/api/resumes", files=docx_file())
    assert uploaded.status_code == 201
    resume_id = uploaded.json()["id"]
    client.post("/api/auth/logout")
    register(client, "b@example.test", "用户B")
    response = client.get(f"/api/resumes/{resume_id}")
    assert response.status_code == 404
    assert response.json()["error"]["code"] == "RESUME_NOT_FOUND"


def test_unsupported_resume_has_clear_error(client):
    register(client)
    response = client.post("/api/resumes", files={"file": ("resume.txt", b"hello", "text/plain")})
    assert response.status_code == 415
    assert "DOCX" in response.json()["error"]["message"]


def test_docx_table_content_is_parsed(client):
    register(client)
    doc = Document()
    doc.add_paragraph("机械工程本科")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "项目经历"
    table.cell(0, 1).text = "使用 CAD 完成机构设计"
    stream = io.BytesIO()
    doc.save(stream)
    response = client.post(
        "/api/resumes",
        files={
            "file": (
                "table-resume.docx",
                stream.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert response.status_code == 201
    detail = client.get(f"/api/resumes/{response.json()['id']}").json()
    assert "项目经历" in detail["parsed_text"] and "CAD" in detail["parsed_text"]


def test_complete_timeline_is_extracted_into_structured_fact_fields(client):
    register(client)
    response = client.post(
        "/api/resumes",
        files=docx_file(
            "姓名：张三\n邮箱：zhangsan@example.test\n工作经历\n"
            "示例制造有限公司 | 机械设计实习生 | 杭州 | 2024-03 至 2024-09\n"
            "负责 CAD 工程图维护"
        ),
    )
    assert response.status_code == 201, response.text
    resume = response.json()
    facts = client.get(f"/api/resumes/{resume['id']}/facts").json()["items"]
    employment = next(fact for fact in facts if fact["category"] == "employment")
    assert employment["organization"] == "示例制造有限公司"
    assert employment["role_title"] == "机械设计实习生"
    assert employment["location"] == "杭州"
    assert employment["start_date"] == "2024-03"
    assert employment["end_date"] == "2024-09"
    detail = client.get(f"/api/resumes/{resume['id']}").json()
    assert detail["contact_json"]["email"] == "zhangsan@example.test"


def test_bulk_confirmation_blocks_missing_critical_fields_then_accepts_answers(client):
    register(client)
    resume = client.post(
        "/api/resumes",
        files=docx_file("项目经历\n使用 CAD 完成机构设计"),
    ).json()
    blocked = client.post(f"/api/resumes/{resume['id']}/confirm-facts")
    assert blocked.status_code == 409
    assert blocked.json()["error"]["code"] == "RESUME_FACTS_INCOMPLETE"
    missing = blocked.json()["error"]["details"]["missing_information"]
    assert {item["field"] for item in missing} == {"organization", "start_date"}
    fact_id = missing[0]["fact_id"]

    answered = client.post(
        f"/api/resumes/{resume['id']}/answer-clarifications",
        json=[
            {"fact_id": fact_id, "field": "organization", "value": "个人项目"},
            {"fact_id": fact_id, "field": "start_date", "value": "2024-03"},
        ],
    )
    assert answered.status_code == 200, answered.text
    assert answered.json()["missing_information"] == []
    confirmed = client.post(f"/api/resumes/{resume['id']}/confirm-facts")
    assert confirmed.status_code == 200, confirmed.text


def test_unverified_dates_from_extractor_are_not_accepted():
    source = "项目经历\n个人完成 CAD 机构设计"
    value = {
        "contact": {},
        "facts": [
            {
                "category": "projects",
                "label": "个人完成 CAD 机构设计",
                "organization": "",
                "role_title": "",
                "location": "",
                "start_date": "2024-03",
                "end_date": "2024-06",
                "text": "个人完成 CAD 机构设计",
                "source_text": "个人完成 CAD 机构设计",
            }
        ],
        "missing_information": [],
    }
    fact = sanitize_extraction(value, source)["facts"][0]
    assert fact["start_date"] == ""
    assert fact["end_date"] == ""


def test_text_pdf_is_uploaded_and_parsed(client):
    register(client)
    response = client.post("/api/resumes", files=pdf_file())
    assert response.status_code == 201, response.text
    detail = client.get(f"/api/resumes/{response.json()['id']}").json()
    assert "Mechanical engineering" in detail["parsed_text"]
    assert client.get(f"/api/resumes/{response.json()['id']}/facts").json()["total"] >= 1


def test_job_normalizer_is_deterministic():
    jd = "设备工程师\n负责设备点检和 SOP 编制\n要求机械相关本科，接受轮班，熟悉 CAD"
    assert normalize_job(jd, company_name="示例公司") == normalize_job(jd, company_name="示例公司")
    result = normalize_job(jd)
    assert "设备点检" in result["keywords"] and "轮班" in result["work_conditions"][0]


def test_matching_boundaries_and_hard_block(db):
    class Fact:
        id = "f1"
        verification_status = "user_confirmed"
        value_json = {"text": "机械工程本科，使用 CAD 完成项目设计"}

    job = Job(
        created_by="u",
        company_name="c",
        title="t",
        raw_jd="x",
        keywords_json=["机械", "本科", "CAD"],
        work_conditions_json=["要求轮班"],
        content_hash="h",
        published_at=datetime.now(timezone.utc),
        status="active",
    )
    profile = UserProfile(user_id="u", accepts_shift=False)
    result = score_match([Fact()], job, profile)
    assert result["hard_block"] is True
    assert result["recommendation"] in {"低优先", "不建议"}
    job.published_at = datetime.now(timezone.utc) - timedelta(days=45)
    assert (
        score_match([Fact()], job, UserProfile(user_id="u", accepts_shift=True))[
            "dimension_scores"
        ]["freshness"]
        == 1
    )


def test_matching_evidence_only_references_confirmed_facts():
    confirmed = type(
        "Fact",
        (),
        {
            "id": "confirmed",
            "verification_status": "user_confirmed",
            "value_json": {"text": "使用 CAD 完成机械项目设计"},
        },
    )()
    unconfirmed = type(
        "Fact",
        (),
        {
            "id": "unconfirmed",
            "verification_status": "extracted",
            "value_json": {"text": "使用 CAD 完成机械项目设计"},
        },
    )()
    job = Job(
        created_by="u",
        company_name="c",
        title="t",
        raw_jd="x",
        keywords_json=["CAD"],
        content_hash="confirmed-evidence",
    )
    result = score_match([unconfirmed, confirmed], job)
    assert result["evidence"][0]["fact_ids"] == ["confirmed"]


def test_manual_job_keeps_unknown_publication_date(client):
    register(client)
    resume = client.post("/api/resumes", files=docx_file()).json()
    client.post(f"/api/resumes/{resume['id']}/confirm-facts")
    job = client.post(
        "/api/jobs/import-text",
        json={
            "company_name": "虚构公司",
            "title": "设备工程师",
            "raw_jd": "负责设备点检和维护，要求机械相关本科并熟悉 CAD。",
        },
    ).json()
    assert job["published_at"] is None
    match = client.post(
        "/api/matches",
        json={"resume_id": resume["id"], "job_id": job["id"]},
    ).json()
    assert match["dimension_scores_json"]["freshness"] == 2


def test_fact_edit_and_new_fact_require_confirmation_again(client):
    register(client)
    resume = client.post("/api/resumes", files=docx_file()).json()
    facts = client.get(f"/api/resumes/{resume['id']}/facts").json()["items"]
    client.post(f"/api/resumes/{resume['id']}/confirm-facts")
    assert client.get(f"/api/resumes/{resume['id']}").json()["facts_confirmed_at"]

    edited = client.patch(
        f"/api/resume-facts/{facts[0]['id']}",
        json={"value_json": {"text": "用户修正后的真实内容"}},
    )
    assert edited.status_code == 200
    assert client.get(f"/api/resumes/{resume['id']}").json()["facts_confirmed_at"] is None
    client.post(f"/api/resume-facts/{facts[0]['id']}/confirm")
    assert client.get(f"/api/resumes/{resume['id']}").json()["facts_confirmed_at"]

    created = client.post(
        f"/api/resumes/{resume['id']}/facts",
        json={
            "category": "projects",
            "label": "个人练习项目",
            "value_json": {"text": "完成虚构设备点检表练习项目"},
            "source_text": "用户主动补充",
        },
    )
    assert created.status_code == 201
    assert created.json()["verification_status"] == "proposed"
    assert client.get(f"/api/resumes/{resume['id']}").json()["facts_confirmed_at"] is None
    client.post(f"/api/resume-facts/{created.json()['id']}/confirm")
    assert client.get(f"/api/resumes/{resume['id']}").json()["facts_confirmed_at"]


def test_matching_api_shows_dimensions_evidence_gaps_and_confirmations(client):
    register(client)
    resume = client.post("/api/resumes", files=docx_file()).json()
    client.post(f"/api/resumes/{resume['id']}/confirm-facts")
    job = client.post(
        "/api/jobs/import-text",
        json={
            "company_name": "示例公司",
            "title": "设备工程师",
            "raw_jd": "负责设备点检和 SOP 编制。要求机械相关本科，接受轮班，熟悉 CAD。",
        },
    ).json()
    response = client.post("/api/matches", json={"resume_id": resume["id"], "job_id": job["id"]})
    assert response.status_code == 201, response.text
    data = response.json()
    assert len(data["dimension_scores_json"]) == 6
    assert isinstance(data["evidence_json"], list) and isinstance(data["gaps_json"], list)
    assert "是否接受轮班" in data["confirmations_json"]
    assert data["ai_explanation_json"]["provider"] == "mock"
