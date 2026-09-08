import io

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from .resume_style import (
    ACCENT_HEX,
    BODY_CJK_FONT,
    BODY_LATIN_FONT,
    MUTED_HEX,
    blueprint_footer,
    format_entry_headline,
    header_copy,
    is_education_section,
    is_skill_tag_section,
    ordered_visible_sections,
    resolved_style,
    should_mark_unachieved,
    unachieved_mark,
)

MUTED = RGBColor.from_string(MUTED_HEX)


def _set_style_font(
    style,
    *,
    size: float,
    color: RGBColor,
    bold: bool = False,
    cjk_font: str = BODY_CJK_FONT,
    latin_font: str = BODY_LATIN_FONT,
) -> None:
    style.font.name = latin_font
    style.font.size = Pt(size)
    style.font.color.rgb = color
    style.font.bold = bold
    style._element.rPr.rFonts.set(qn("w:ascii"), latin_font)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), latin_font)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), cjk_font)


def _snapshot_style(snapshot: dict | None) -> dict:
    values = resolved_style(snapshot)
    return {
        **values,
        "ink": RGBColor.from_string(values["ink_hex"]),
        "accent": RGBColor.from_string(values["accent_hex"]),
        "muted": RGBColor.from_string(values["muted_hex"]),
    }


def render_resume_docx(content: dict, fallback_title: str, snapshot: dict | None = None) -> bytes:
    """Render v1 or v2 resume content as a deterministic A4 Word document."""
    style = _snapshot_style(snapshot)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(style["page_width_mm"])
    section.page_height = Mm(style["page_height_mm"])
    section.top_margin = Mm(style["margin_vertical_mm"])
    section.right_margin = Mm(style["margin_horizontal_mm"])
    section.bottom_margin = Mm(style["margin_vertical_mm"])
    section.left_margin = Mm(style["margin_horizontal_mm"])
    section.header_distance = Mm(8)
    section.footer_distance = Mm(8)

    _configure_styles(doc, style)
    header = header_copy(content, fallback_title)
    title = doc.add_paragraph(style="Resume Name")
    title.add_run(header["display_name"])
    if header["subtitle"]:
        subtitle = doc.add_paragraph(style="Resume Subtitle")
        subtitle.add_run(header["subtitle"])
    if header["contact_parts"]:
        paragraph = doc.add_paragraph(style="Resume Contact")
        paragraph.add_run("    ".join(header["contact_parts"]))

    gap = content.get("gap_summary") if isinstance(content.get("gap_summary"), dict) else {}
    if content.get("kind") == "target" and gap.get("reasons"):
        doc.add_paragraph(str(gap.get("title") or "对照岗位还缺什么"), style="Resume Section")
        intro = doc.add_paragraph(style="Resume Contact")
        intro.add_run("未达成条目按 JD 列出，不是已经做过的经历。")
        for reason in gap.get("reasons") or []:
            _add_bullet(doc, str(reason), achieved=True, style=style)

    for value in ordered_visible_sections(content, snapshot):
        items = value.get("items") if isinstance(value.get("items"), list) else []
        if not items:
            continue
        heading = doc.add_paragraph(str(value.get("name") or "经历"), style="Resume Section")
        _set_bottom_border(heading, accent_hex=style["accent_hex"])
        if is_skill_tag_section(value):
            _render_skill_tags(
                doc, items, mark=unachieved_mark(content), content=content, style=style
            )
            continue
        education = is_education_section(value)
        for item in items:
            if not isinstance(item, dict):
                continue
            if isinstance(item.get("bullets"), list):
                _render_v2_item(
                    doc,
                    item,
                    education=education,
                    mark=unachieved_mark(content),
                    content=content,
                    style=style,
                )
            elif item.get("text"):
                _add_bullet(doc, str(item.get("text") or ""), achieved=True, mark="", style=style)

    footer = blueprint_footer(content)
    if footer:
        note = doc.add_paragraph(style="Resume Contact")
        note.add_run(footer)

    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


def _configure_styles(doc: Document, style: dict) -> None:
    normal = doc.styles["Normal"]
    _set_style_font(
        normal,
        size=style["body_size_pt"],
        color=style["ink"],
        cjk_font=style["cjk_font"],
        latin_font=style["latin_font"],
    )
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = style["line_spacing"]

    name = doc.styles.add_style("Resume Name", WD_STYLE_TYPE.PARAGRAPH)
    _set_style_font(
        name,
        size=style["name_size_pt"],
        color=style["accent"],
        bold=True,
        cjk_font=style["heading_cjk"],
        latin_font=style["latin_font"],
    )
    name.paragraph_format.space_after = Pt(3)
    name.paragraph_format.keep_with_next = True

    subtitle = doc.styles.add_style("Resume Subtitle", WD_STYLE_TYPE.PARAGRAPH)
    _set_style_font(
        subtitle,
        size=style["section_size_pt"],
        color=style["accent"],
        bold=True,
        cjk_font=style["heading_cjk"],
        latin_font=style["latin_font"],
    )
    subtitle.paragraph_format.space_after = Pt(2)

    contact = doc.styles.add_style("Resume Contact", WD_STYLE_TYPE.PARAGRAPH)
    _set_style_font(
        contact,
        size=9.5,
        color=style["ink"],
        cjk_font=style["cjk_font"],
        latin_font=style["latin_font"],
    )
    contact.paragraph_format.space_after = Pt(4)
    contact.paragraph_format.keep_with_next = True

    heading = doc.styles.add_style("Resume Section", WD_STYLE_TYPE.PARAGRAPH)
    _set_style_font(
        heading,
        size=style["section_size_pt"],
        color=style["accent"],
        bold=True,
        cjk_font=style["heading_cjk"],
        latin_font=style["latin_font"],
    )
    heading.paragraph_format.space_before = Pt(7)
    heading.paragraph_format.space_after = Pt(3)
    heading.paragraph_format.keep_with_next = True

    entry = doc.styles.add_style("Resume Entry Header", WD_STYLE_TYPE.PARAGRAPH)
    _set_style_font(
        entry,
        size=style["body_size_pt"],
        color=style["ink"],
        bold=True,
        cjk_font=style["cjk_font"],
        latin_font=style["latin_font"],
    )
    entry.paragraph_format.space_before = Pt(2)
    entry.paragraph_format.space_after = Pt(1)
    entry.paragraph_format.keep_with_next = True
    entry.paragraph_format.tab_stops.add_tab_stop(Mm(173), WD_TAB_ALIGNMENT.RIGHT)

    bullet = doc.styles.add_style("Resume Bullet", WD_STYLE_TYPE.PARAGRAPH)
    bullet.base_style = doc.styles["List Bullet"]
    _set_style_font(
        bullet,
        size=style["body_size_pt"],
        color=style["ink"],
        cjk_font=style["cjk_font"],
        latin_font=style["latin_font"],
    )
    bullet.paragraph_format.left_indent = Mm(5)
    bullet.paragraph_format.first_line_indent = Mm(-3)
    bullet.paragraph_format.space_after = Pt(2)
    bullet.paragraph_format.line_spacing = style["line_spacing"]
    bullet.paragraph_format.keep_together = True

    skills = doc.styles.add_style("Resume Skills", WD_STYLE_TYPE.PARAGRAPH)
    _set_style_font(
        skills,
        size=style["body_size_pt"],
        color=style["ink"],
        cjk_font=style["cjk_font"],
        latin_font=style["latin_font"],
    )
    skills.paragraph_format.space_after = Pt(3)
    skills.paragraph_format.line_spacing = style["line_spacing"]


def _render_skill_tags(
    doc: Document,
    items: list,
    *,
    mark: str,
    content: dict | None = None,
    style: dict | None = None,
) -> None:
    parts = []
    muted_indexes = set()
    for item in items:
        if not isinstance(item, dict):
            continue
        for bullet in item.get("bullets") or []:
            if not isinstance(bullet, dict):
                continue
            text = str(bullet.get("text") or "").strip()
            if not text:
                continue
            if should_mark_unachieved(content, bullet):
                muted_indexes.add(len(parts))
                text = f"{text}{mark}"
            parts.append(text)
    if not parts:
        return
    paragraph = doc.add_paragraph(style="Resume Skills")
    for index, part in enumerate(parts):
        if index:
            paragraph.add_run(" / ")
        run = paragraph.add_run(part)
        if index in muted_indexes:
            run.font.color.rgb = (style or {}).get("muted") or MUTED


def _set_bottom_border(paragraph, *, accent_hex: str = ACCENT_HEX) -> None:
    paragraph_properties = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), accent_hex)
    borders.append(bottom)
    paragraph_properties.append(borders)


def _render_v2_item(
    doc: Document,
    item: dict,
    *,
    education: bool = False,
    mark: str,
    content: dict | None = None,
    style: dict | None = None,
) -> None:
    headline = format_entry_headline(item, education=education)
    if headline:
        header = doc.add_paragraph(style="Resume Entry Header")
        header.add_run(headline)
    for bullet in item.get("bullets") or []:
        if not isinstance(bullet, dict):
            continue
        text = str(bullet.get("text") or "").strip()
        if not text:
            continue
        _add_bullet(
            doc,
            text,
            achieved=not should_mark_unachieved(content, bullet),
            mark=mark if should_mark_unachieved(content, bullet) else "",
            style=style,
        )


def _add_bullet(
    doc: Document, text: str, *, achieved: bool, mark: str = "", style: dict | None = None
) -> None:
    value = str(text or "").strip()
    if not value:
        return
    muted = (style or {}).get("muted") or RGBColor.from_string(MUTED_HEX)
    paragraph = doc.add_paragraph(style="Resume Bullet")
    if not achieved:
        run = paragraph.add_run(value)
        run.font.color.rgb = muted
        if mark:
            label = paragraph.add_run(mark)
            label.font.superscript = True
            label.font.color.rgb = muted
    else:
        paragraph.add_run(value)
