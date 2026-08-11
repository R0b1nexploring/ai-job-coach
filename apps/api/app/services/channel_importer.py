import csv
import hashlib
import json
import re
import unicodedata
from collections import Counter
from datetime import UTC, date, datetime
from pathlib import Path
from urllib.parse import urlsplit

from email_validator import EmailNotValidError, validate_email
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.entities import ApplicationChannel
from app.providers.jobs.base import JobProviderError
from app.providers.jobs.public_jsonld import _ensure_public_addresses, _resolve_host

from .channel_sources import invalidate_channel_host_cache

FIELD_ALIASES = {
    "company_name": ("company_name", "company", "公司", "公司名称"),
    "company_aliases": ("company_aliases", "aliases", "公司别名"),
    "channel_type": ("channel_type", "type", "渠道类型"),
    "url": ("url", "网址", "投递网址", "招聘页"),
    "email": ("email", "邮箱", "投递邮箱"),
    "scope": ("scope", "招聘范围", "招聘类型"),
    "regions": ("regions", "地区", "区域"),
    "industries": ("industries", "行业"),
    "source_updated_at": ("source_updated_at", "updated_at", "更新时间"),
    "notes": ("notes", "备注"),
}
FORBIDDEN_HOSTS = {"localhost"}
FORBIDDEN_SUFFIXES = (".local", ".internal", ".localdomain")


class ChannelImportError(ValueError):
    def __init__(self, reason: str):
        super().__init__(reason)
        self.reason = reason


def normalize_company_name(value: str) -> str:
    normalized = unicodedata.normalize("NFKC", value or "").casefold()
    return re.sub(r"[\s·•\-_—–()（）]+", "", normalized)


def read_channel_rows(path: Path) -> list[dict]:
    suffix = path.suffix.casefold()
    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as stream:
            return list(csv.DictReader(stream))
    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        values = sheet.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(values, ())]
        return [dict(zip(headers, row, strict=False)) for row in values]
    raise ChannelImportError("unsupported_file_type")


def import_channels(
    db: Session,
    rows: list[dict],
    *,
    source: str,
    execute: bool = False,
    owner_user_id: str | None = None,
    column_map: dict[str, str] | None = None,
    resolver=None,
    default_source_updated_at: datetime | None = None,
) -> dict:
    stats = {"new": 0, "updated": 0, "skipped": 0, "rejected": 0}
    reasons: Counter[str] = Counter()
    seen_hashes: set[str] = set()
    for raw in rows:
        try:
            values = normalize_channel_row(
                raw,
                source=source,
                owner_user_id=owner_user_id,
                column_map=column_map,
                resolver=resolver,
                default_source_updated_at=default_source_updated_at,
            )
        except ChannelImportError as exc:
            stats["rejected"] += 1
            reasons[exc.reason] += 1
            continue
        content_hash = values["content_hash"]
        if content_hash in seen_hashes:
            stats["skipped"] += 1
            continue
        seen_hashes.add(content_hash)
        existing = db.scalar(
            select(ApplicationChannel).where(ApplicationChannel.content_hash == content_hash)
        )
        if existing:
            stats["updated"] += 1
            if execute:
                existing.source_updated_at = values["source_updated_at"]
                existing.source = values["source"]
                existing.notes = values["notes"]
            continue
        stats["new"] += 1
        if execute:
            db.add(ApplicationChannel(**values))
    if execute:
        db.commit()
        invalidate_channel_host_cache(db)
    return {**stats, "rejected_reasons": dict(sorted(reasons.items())), "executed": execute}


def normalize_channel_row(
    raw: dict,
    *,
    source: str,
    owner_user_id: str | None = None,
    column_map: dict[str, str] | None = None,
    resolver=None,
    default_source_updated_at: datetime | None = None,
) -> dict:
    def value(field: str) -> str:
        return _field(raw, field, column_map).strip()

    company_name = value("company_name")
    company_key = normalize_company_name(company_name)
    if not company_key:
        raise ChannelImportError("company_required")
    channel_type = value("channel_type").casefold()
    if channel_type not in {"career_page", "email", "form"}:
        raise ChannelImportError("invalid_channel_type")
    url = value("url")
    email = value("email")
    if channel_type in {"career_page", "form"} and not url:
        raise ChannelImportError("url_required")
    if channel_type == "email" and not email:
        raise ChannelImportError("email_required")
    if url:
        _validate_import_url(url, resolver=resolver)
    if email:
        try:
            email = validate_email(email, check_deliverability=False).normalized
        except EmailNotValidError as exc:
            raise ChannelImportError("invalid_email") from exc
    scope = value("scope").casefold() or "both"
    if scope not in {"campus", "experienced", "both"}:
        raise ChannelImportError("invalid_scope")
    source_updated_at = _parse_datetime(value("source_updated_at"))
    if source_updated_at is None:
        source_updated_at = default_source_updated_at
    aliases = _list_value(value("company_aliases"))
    regions = _list_value(value("regions"))
    industries = _list_value(value("industries"))
    identity = json.dumps(
        [owner_user_id or "global", company_key, channel_type, url, email],
        ensure_ascii=False,
        separators=(",", ":"),
    )
    return {
        "owner_user_id": owner_user_id,
        "company_name": company_name[:200],
        "company_name_normalized": company_key[:200],
        "company_aliases_json": aliases,
        "channel_type": channel_type,
        "url": url[:1000],
        "email": email[:320],
        "scope": scope,
        "regions_json": regions,
        "industries_json": industries,
        "source": source[:200],
        "source_updated_at": source_updated_at,
        "verify_status": "unverified",
        "dead_reports": 0,
        "content_hash": hashlib.sha256(identity.encode()).hexdigest(),
        "notes": value("notes")[:5000],
    }


def _validate_import_url(url: str, *, resolver=None) -> str:
    try:
        parsed = urlsplit(url)
        port = parsed.port
    except ValueError as exc:
        raise ChannelImportError("invalid_url") from exc
    if parsed.scheme.casefold() != "https":
        raise ChannelImportError("url_requires_https")
    if not parsed.hostname:
        raise ChannelImportError("invalid_url")
    if parsed.username or parsed.password:
        raise ChannelImportError("url_userinfo_forbidden")
    if port not in {None, 443}:
        raise ChannelImportError("url_port_forbidden")
    host = parsed.hostname.rstrip(".").casefold()
    if host in FORBIDDEN_HOSTS or host.endswith(FORBIDDEN_SUFFIXES):
        raise ChannelImportError("url_local_host_forbidden")
    try:
        import ipaddress

        ipaddress.ip_address(host)
    except ValueError:
        pass
    else:
        raise ChannelImportError("url_ip_literal_forbidden")
    try:
        addresses = (resolver or _resolve_host)(host)
        _ensure_public_addresses(addresses)
    except JobProviderError as exc:
        raise ChannelImportError("url_not_public") from exc
    return host


def _field(raw: dict, name: str, column_map: dict[str, str] | None) -> str:
    if column_map and column_map.get(name):
        return str(raw.get(column_map[name]) or "")
    lowered = {str(key).strip().casefold(): value for key, value in raw.items()}
    for alias in FIELD_ALIASES[name]:
        if alias.casefold() in lowered:
            return str(lowered[alias.casefold()] or "")
    return ""


def _list_value(value: str) -> list[str]:
    return list(
        dict.fromkeys(
            item.strip() for item in re.split(r"[,，;；|\n]", value or "") if item.strip()
        )
    )[:50]


def _parse_datetime(value: str) -> datetime | None:
    if not value:
        return None
    try:
        parsed = datetime.fromisoformat(value)
    except ValueError:
        try:
            parsed = datetime.combine(date.fromisoformat(value), datetime.min.time())
        except ValueError as exc:
            raise ChannelImportError("invalid_source_updated_at") from exc
    return parsed if parsed.tzinfo else parsed.replace(tzinfo=UTC)
