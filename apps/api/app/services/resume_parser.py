import io
import re
from collections import Counter

import pdfplumber
from docx import Document

from app.errors import AppError

ALLOWED = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
    "application/pdf": ".pdf",
}

FACT_CATEGORIES = {"education", "employment", "projects", "skills", "certificates", "other"}
SECTION_CATEGORIES = {
    "教育背景": "education",
    "教育经历": "education",
    "工作经历": "employment",
    "实习经历": "employment",
    "实习工作": "employment",
    "项目经历": "projects",
    "项目经验": "projects",
    "项目科研竞赛": "projects",
    "技能": "skills",
    "专业技能": "skills",
    "技能证书": "certificates",
    "证书": "certificates",
    "证书资质": "certificates",
    "奖项证书": "certificates",
    "奖项": "certificates",
}
DUTY_PREFIX = re.compile(r"^(公司需要|公司将|负责|通过分析|提升|针对)")
CJK_PAIR_SPACE = re.compile(r"(?<=[\u4e00-\u9fff])[ \t]+(?=[\u4e00-\u9fff])")
DATE_PART = r"(?:19|20)\d{2}(?:[./-](?:0?[1-9]|1[0-2])|年(?:0?[1-9]|1[0-2])月?)"
DATE_RANGE = re.compile(
    rf"(?P<start>{DATE_PART})\s*(?:至|到|[-–—~～])\s*(?P<end>{DATE_PART}|至今|现在)"
)
EMAIL = re.compile(r"(?<![\w.+-])[\w.+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}(?![\w.-])")
PHONE = re.compile(r"(?<!\d)(?:\+?86[- ]?)?1[3-9]\d{9}(?!\d)")
ORGANIZATION_WORDS = ("大学", "学院", "公司", "集团", "实验室", "研究院", "中心", "医院")
ROLE_WORDS = ("工程师", "实习生", "助理", "经理", "负责人", "研究员", "设计师", "专员")
KNOWN_CITIES = {
    "北京",
    "上海",
    "广州",
    "深圳",
    "杭州",
    "南京",
    "苏州",
    "成都",
    "武汉",
    "西安",
    "天津",
    "重庆",
}


def parse_resume(content: bytes, media_type: str, filename: str) -> tuple[str, list[dict]]:
    suffix = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    expected = ALLOWED.get(media_type)
    if not expected or suffix != expected:
        raise AppError("UNSUPPORTED_RESUME_TYPE", "仅支持 DOCX 或文本型 PDF 文件。", 415)
    try:
        if suffix == ".docx":
            doc = Document(io.BytesIO(content))
            parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        parts.append(" | ".join(dict.fromkeys(cells)))
            text = "\n".join(parts)
        else:
            with pdfplumber.open(io.BytesIO(content)) as pdf:
                text = "\n".join((page.extract_text() or "") for page in pdf.pages).strip()
            if not text:
                raise AppError(
                    "SCANNED_PDF_UNSUPPORTED", "该 PDF 可能是扫描件，MVP 暂不支持 OCR。", 422
                )
    except AppError:
        raise
    except Exception as exc:
        raise AppError("RESUME_PARSE_FAILED", "文件无法解析，请确认文件未损坏。", 422) from exc
    if not text.strip():
        raise AppError("EMPTY_RESUME", "简历中没有可解析的文本。", 422)
    text = normalize_extracted_text(text)
    return text, extract_facts(text)


def normalize_extracted_text(text: str) -> str:
    """Collapse PDF-extracted spaces between CJK characters without rewriting wording."""
    collapsed = CJK_PAIR_SPACE.sub("", str(text or ""))
    return re.sub(r"[ \t]+", " ", collapsed)


def preprocess_resume_lines(text: str) -> list[str]:
    """Remove repeated page furniture while preserving source wording and line order."""
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.splitlines()]
    lines = [line for line in lines if len(line) >= 2]
    counts = Counter(lines)
    return [line for line in lines if counts[line] < 3]


def deterministic_extraction(text: str) -> dict:
    facts = extract_facts(text)
    return {
        "contact": extract_contact(text),
        "facts": [
            {
                "category": fact["category"],
                "label": fact["label"],
                "organization": fact["organization"],
                "role_title": fact["role_title"],
                "location": fact["location"],
                "start_date": fact["start_date"],
                "end_date": fact["end_date"],
                "text": fact["value_json"]["text"],
                "source_text": fact["source_text"],
            }
            for fact in facts
        ],
        "missing_information": build_missing_information(facts),
    }


def extract_facts(text: str) -> list[dict]:
    facts = []
    for category, source_text in _fact_blocks(preprocess_resume_lines(text))[:30]:
        metadata = _metadata(source_text)
        facts.append(
            {
                "category": category,
                "label": _label(source_text, metadata),
                "value_json": {"text": source_text},
                "organization": metadata["organization"],
                "role_title": metadata["role_title"],
                "location": metadata["location"],
                "start_date": metadata["start_date"],
                "end_date": metadata["end_date"],
                "source_text": source_text,
            }
        )
    return facts


def extract_contact(text: str) -> dict:
    email = EMAIL.search(text)
    phone = PHONE.search(text)
    name_match = re.search(r"(?:姓名|Name)\s*[:：]\s*([^\n|｜]{2,30})", text, re.IGNORECASE)
    city_match = re.search(r"(?:城市|所在地|现居地)\s*[:：]\s*([^\n|｜]{2,30})", text)
    name = name_match.group(1).strip() if name_match else ""
    if not name:
        first = next((line.strip() for line in text.splitlines() if line.strip()), "")
        inferred = re.match(r"^([\u4e00-\u9fff]{2,4})(?:性别|男|女|电话|邮箱|$)", first)
        name = inferred.group(1) if inferred else ""
    return {
        "name": name,
        "phone": phone.group(0).strip() if phone else "",
        "email": email.group(0).strip() if email else "",
        "city": city_match.group(1).strip() if city_match else "",
    }


def sanitize_extraction(value: dict, source_text: str) -> dict:
    """Discard generated claims that cannot be traced back to the uploaded text."""
    raw_contact = value.get("contact") if isinstance(value.get("contact"), dict) else {}
    contact = {
        key: _evidenced_text(raw_contact.get(key), source_text)
        for key in ("name", "phone", "email", "city")
    }
    facts = []
    for raw in value.get("facts") or []:
        if not isinstance(raw, dict):
            continue
        text_value = str(raw.get("text") or "").strip()
        source_value = str(raw.get("source_text") or "").strip()
        if text_value not in source_text:
            text_value = source_value if source_value and source_value in source_text else ""
        if not text_value:
            continue
        if not source_value or source_value not in source_text:
            source_value = text_value
        category = str(raw.get("category") or "other").strip().lower()
        category = category if category in FACT_CATEGORIES else "other"
        start_date = _evidenced_date(raw.get("start_date"), source_value)
        end_raw = str(raw.get("end_date") or "").strip()
        end_date = end_raw if end_raw in {"至今", "现在"} and end_raw in source_value else ""
        if not end_date:
            end_date = _evidenced_date(end_raw, source_value)
        fact = {
            "category": category,
            "label": _evidenced_label(raw.get("label"), text_value),
            "value_json": {"text": text_value},
            "organization": _evidenced_text(raw.get("organization"), source_value),
            "role_title": _evidenced_text(raw.get("role_title"), source_value),
            "location": _evidenced_text(raw.get("location"), source_value),
            "start_date": start_date,
            "end_date": "至今" if end_date == "现在" else end_date,
            "source_text": source_value,
        }
        facts.append(fact)
    return {
        "contact": contact,
        "facts": facts,
        "missing_information": build_missing_information(facts),
    }


def build_missing_information(facts: list) -> list[dict]:
    missing = []
    for index, fact in enumerate(facts):
        category = _fact_value(fact, "category")
        if category not in {"employment", "projects"}:
            continue
        fact_id = _fact_value(fact, "id") or None
        label = _fact_value(fact, "label") or f"第 {index + 1} 条经历"
        for field, field_label in (("organization", "组织名称"), ("start_date", "开始时间")):
            if _fact_value(fact, field):
                continue
            missing.append(
                {
                    "fact_index": index,
                    "fact_id": fact_id,
                    "field": field,
                    "question": f"“{label}”的{field_label}是什么？",
                }
            )
    return missing


def _fact_blocks(lines: list[str]) -> list[tuple[str, str]]:
    blocks: list[tuple[str, list[str]]] = []
    section = "other"
    current_category = ""
    current_lines: list[str] = []

    def flush() -> None:
        nonlocal current_category, current_lines
        if current_lines:
            blocks.append((current_category or "other", list(current_lines)))
        current_category = ""
        current_lines = []

    for line in lines:
        heading = re.sub(r"[：:\s]", "", line)
        heading_category = _heading_category(heading)
        if heading_category:
            flush()
            section = heading_category
            continue
        inferred = _infer_category(line, section)
        if current_lines and (
            (inferred != "other" and inferred != current_category) or _starts_new_fact(line)
        ):
            flush()
        if not current_lines:
            current_category = inferred
        current_lines.append(line)
    flush()
    return [(category, "\n".join(parts)) for category, parts in blocks]


def _infer_category(line: str, section: str) -> str:
    patterns = (
        ("education", r"大学|学院|本科|硕士|博士|专业|学历"),
        ("certificates", r"证书|资格证|CET|PMP|CFA|六西格玛"),
        ("skills", r"技能|CAD|SolidWorks|Python|SOP|Excel|PLC"),
        ("projects", r"项目|设计|课题|作品"),
        ("employment", r"实习|工作|任职|公司|工程师"),
    )
    if section != "other":
        return section
    return next(
        (category for category, pattern in patterns if re.search(pattern, line, re.IGNORECASE)),
        "other",
    )


def _heading_category(heading: str) -> str | None:
    if heading in SECTION_CATEGORIES:
        return SECTION_CATEGORIES[heading]
    for name, category in sorted(SECTION_CATEGORIES.items(), key=lambda item: -len(item[0])):
        if heading.startswith(name):
            return category
    return None


def _starts_new_fact(line: str) -> bool:
    if DATE_RANGE.search(line):
        return True
    if DUTY_PREFIX.match(line) or len(line) > 40:
        return False
    return any(word in line for word in ORGANIZATION_WORDS)


def _metadata(source_text: str) -> dict:
    first_line = source_text.splitlines()[0]
    segments = [part.strip() for part in re.split(r"[|｜·•]", first_line) if part.strip()]
    date_match = DATE_RANGE.search(source_text)
    organization = next((part for part in segments if _looks_like_organization(part)), "")
    if not organization:
        organization = next(
            (
                part
                for part in segments
                if 2 <= len(part) <= 16
                and not DATE_RANGE.search(part)
                and not DUTY_PREFIX.match(part)
                and not re.search(r"使用|完成|进行|负责|参与", part)
                and not any(word in part for word in ROLE_WORDS)
            ),
            "",
        )
    role_title = next((part for part in segments if any(word in part for word in ROLE_WORDS)), "")
    location = next(
        (
            part
            for part in segments
            if part in KNOWN_CITIES or (len(part) <= 12 and part.endswith(("市", "省")))
        ),
        "",
    )
    return {
        "organization": organization,
        "role_title": role_title,
        "location": location,
        "start_date": _normalize_date(date_match.group("start")) if date_match else "",
        "end_date": _normalize_date(date_match.group("end")) if date_match else "",
    }


def _normalize_date(value: str) -> str:
    raw = str(value or "").strip()
    if raw in {"至今", "现在"}:
        return "至今"
    match = re.fullmatch(r"((?:19|20)\d{2})[./-](0?[1-9]|1[0-2])", raw)
    if not match:
        match = re.fullmatch(r"((?:19|20)\d{2})年(0?[1-9]|1[0-2])月?", raw)
    return f"{match.group(1)}-{int(match.group(2)):02d}" if match else ""


def _evidenced_date(value, source_text: str) -> str:
    normalized = _normalize_date(str(value or ""))
    if not normalized:
        return ""
    evidenced_dates = {
        _normalize_date(match.group(0)) for match in re.finditer(DATE_PART, source_text)
    }
    return normalized if normalized in evidenced_dates else ""


def _looks_like_organization(part: str) -> bool:
    text = str(part or "").strip()
    if not text or len(text) > 40 or DUTY_PREFIX.match(text):
        return False
    return any(word in text for word in ORGANIZATION_WORDS)


def _evidenced_text(value, source_text: str) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    if text.casefold() in source_text.casefold():
        return text
    compact_value = re.sub(r"\s+", "", text).casefold()
    compact_source = re.sub(r"\s+", "", source_text).casefold()
    return text if compact_value and compact_value in compact_source else ""


def _evidenced_label(value, source_text: str) -> str:
    label = _evidenced_text(value, source_text)
    return label[:200] if label else source_text.splitlines()[0][:30]


def _label(source_text: str, metadata: dict) -> str:
    return (
        metadata.get("organization")
        or metadata.get("role_title")
        or source_text.splitlines()[0][:30]
    )


def _fact_value(fact, field: str):
    return fact.get(field, "") if isinstance(fact, dict) else getattr(fact, field, "")
