import io
import uuid
from datetime import datetime, timezone
from typing import Literal

from docx import Document
from fastapi import APIRouter, Depends, Header, Query, Response
from fastapi.responses import HTMLResponse, JSONResponse
from pydantic import BaseModel
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.database import get_db
from app.dependencies import current_user
from app.errors import AppError
from app.models.entities import (
    Application,
    ApplicationEvent,
    AuditLog,
    Interview,
    Job,
    Match,
    Notification,
    Order,
    Package,
    QuotaLedger,
    Resume,
    ResumeFact,
    ResumeVersion,
    User,
    UserProfile,
)
from app.providers.llm.runtime import generate_structured
from app.providers.llm.schemas import InterviewStructuredOutput, ResumeStructuredOutput
from app.providers.storage.local import LocalStorage
from app.services.fact_guard import guard_content
from app.services.quota_service import balance, change
from app.services.report_service import weekly_report
from app.services.resume_generator import generate_resume
from app.services.interview_generator import generate_questions
from app.services.html_resume_renderer import render_resume_html
from .common import dump, listed


router = APIRouter(prefix="/api", tags=["workflow"])


class VersionIn(BaseModel):
    base_resume_id: str
    job_id: str
    match_id: str
    inject_unverified: bool = False


@router.post("/resume-versions", status_code=201)
async def create_version(
    data: VersionIn,
    idempotency_key: str | None = Header(default=None, alias="Idempotency-Key"),
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    resume = db.scalar(
        select(Resume).where(
            Resume.id == data.base_resume_id, Resume.user_id == user.id, Resume.deleted_at.is_(None)
        )
    )
    job = db.scalar(select(Job).where(Job.id == data.job_id, Job.created_by == user.id))
    match = db.scalar(
        select(Match).where(
            Match.id == data.match_id,
            Match.user_id == user.id,
            Match.resume_id == data.base_resume_id,
            Match.job_id == data.job_id,
        )
    )
    if not all([resume, job, match]):
        raise AppError("RESOURCE_NOT_FOUND", "简历、岗位或匹配报告不存在。", 404)
    if not resume.facts_confirmed_at:
        raise AppError("RESUME_FACTS_NOT_CONFIRMED", "请先确认简历事实库后再生成定制版本。", 409)
    key = idempotency_key or f"resume-version:{user.id}:{uuid.uuid4()}"
    existing_ledger = db.scalar(
        select(QuotaLedger).where(
            QuotaLedger.user_id == user.id,
            QuotaLedger.quota_type == "resume",
            QuotaLedger.idempotency_key == key,
        )
    )
    if existing_ledger and existing_ledger.reference_id:
        existing = db.scalar(
            select(ResumeVersion).where(
                ResumeVersion.id == existing_ledger.reference_id,
                ResumeVersion.user_id == user.id,
            )
        )
        if existing:
            return dump(existing)
    if existing_ledger:
        raise AppError("IDEMPOTENCY_KEY_REUSED", "该幂等键已用于其他额度操作，请更换后重试。", 409)
    facts = list(
        db.scalars(
            select(ResumeFact).where(
                ResumeFact.resume_id == resume.id, ResumeFact.user_id == user.id
            )
        )
    )
    fallback_content, _fallback_status, _fallback_claims = generate_resume(
        facts, job, data.inject_unverified
    )
    confirmed_facts = [fact for fact in facts if fact.verification_status == "user_confirmed"]
    structured, provider_name = await generate_structured(
        task="tailored_resume",
        system_prompt=(
            "你是求职简历编辑。只能重写、压缩、排序用户已确认事实；每个要点必须引用 "
            "source_fact_ids。不得新增公司、岗位、日期、工具、证书、成果或数字。"
            "无证据要求放入 excluded_gaps。输出严格 JSON。"
        ),
        user_payload={
            "confirmed_facts": [
                {
                    "id": fact.id,
                    "category": fact.category,
                    "label": fact.label,
                    "value": fact.value_json,
                }
                for fact in confirmed_facts
            ],
            "job": {
                "title": job.title,
                "company_name": job.company_name,
                "must_have": job.must_have_json,
                "keywords": job.keywords_json,
            },
            "mock_response": fallback_content,
        },
        response_model=ResumeStructuredOutput,
    )
    content = structured.model_dump()
    guard_status, new_claims = guard_content(content, confirmed_facts)
    version = ResumeVersion(
        user_id=user.id,
        base_resume_id=resume.id,
        job_id=job.id,
        title=f"{job.company_name} - {job.title}",
        content_json=content,
        traceability_json={"new_claims": new_claims, "provider": provider_name},
        guard_status=guard_status,
    )
    db.add(version)
    db.flush()
    change(
        db,
        user_id=user.id,
        quota_type="resume",
        delta=-1,
        reason="生成定制简历",
        idempotency_key=key,
        reference_type="resume_version",
        reference_id=version.id,
    )
    db.commit()
    return dump(version)


def owned_version(db, user, version_id):
    row = db.scalar(
        select(ResumeVersion).where(
            ResumeVersion.id == version_id, ResumeVersion.user_id == user.id
        )
    )
    if not row:
        raise AppError("RESUME_VERSION_NOT_FOUND", "定制简历不存在。", 404)
    return row


@router.get("/resume-versions")
async def versions(user: User = Depends(current_user), db: Session = Depends(get_db)):
    return listed(
        [
            dump(x)
            for x in db.scalars(
                select(ResumeVersion)
                .where(ResumeVersion.user_id == user.id)
                .order_by(ResumeVersion.created_at.desc())
            ).all()
        ]
    )


@router.get("/resume-versions/{version_id}")
async def version_detail(
    version_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)
):
    return dump(owned_version(db, user, version_id))


@router.post("/resume-versions/{version_id}/submit-review")
async def submit_review(
    version_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)
):
    row = owned_version(db, user, version_id)
    if row.guard_status != "passed":
        raise AppError(
            "FACT_GUARD_BLOCKED",
            "检测到未经确认的新事实，不能提交审核。",
            409,
            row.traceability_json,
        )
    if row.review_status == "approved":
        return dump(row)
    row.review_status = "advisor_review"
    db.add(
        AuditLog(
            actor_user_id=user.id,
            action="resume_version.submitted",
            entity_type="resume_version",
            entity_id=row.id,
        )
    )
    db.commit()
    return dump(row)


def exportable_version(db, user, version_id):
    row = owned_version(db, user, version_id)
    if row.guard_status != "passed" or row.review_status != "approved":
        raise AppError("RESUME_NOT_APPROVED", "仅可导出事实检查通过且已审核的简历。", 409)
    return row


def html_security_headers(disposition: str) -> dict[str, str]:
    return {
        "Content-Disposition": disposition,
        "Content-Security-Policy": "default-src 'none'; style-src 'unsafe-inline'; script-src 'unsafe-inline'; img-src data:",
        "Cache-Control": "private, no-store",
        "X-Content-Type-Options": "nosniff",
    }


def audit_export(db: Session, user: User, row: ResumeVersion, format: str) -> None:
    db.add(
        AuditLog(
            actor_user_id=user.id,
            action="resume_version.exported",
            entity_type="resume_version",
            entity_id=row.id,
            metadata_json={"format": format},
        )
    )
    db.commit()


@router.get("/resume-versions/{version_id}/preview", response_class=HTMLResponse)
async def preview_version(
    version_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)
):
    row = exportable_version(db, user, version_id)
    audit_export(db, user, row, "html_preview")
    return HTMLResponse(
        render_resume_html(row.content_json, row.title), headers=html_security_headers("inline")
    )


@router.get("/resume-versions/{version_id}/download")
async def download_version(
    version_id: str,
    format: Literal["docx", "html"] = Query(default="docx"),
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    row = exportable_version(db, user, version_id)
    if format == "html":
        audit_export(db, user, row, "html_download")
        return HTMLResponse(
            render_resume_html(row.content_json, row.title),
            headers=html_security_headers('attachment; filename="tailored-resume.html"'),
        )
    if not row.export_file_key:
        doc = Document()
        doc.add_heading(row.content_json.get("headline", row.title), 0)
        for section in row.content_json.get("sections", []):
            doc.add_heading(section.get("name", "经历"), level=1)
            for item in section.get("items", []):
                doc.add_paragraph(item.get("text", ""), style="List Bullet")
        stream = io.BytesIO()
        doc.save(stream)
        row.export_file_key = LocalStorage().save(stream.getvalue(), ".docx")
        db.commit()
    content = LocalStorage().path(row.export_file_key).read_bytes()
    audit_export(db, user, row, "docx")
    return Response(
        content=content,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": 'attachment; filename="tailored-resume.docx"'},
    )


class ApplicationIn(BaseModel):
    job_id: str
    resume_version_id: str | None = None
    notes: str = ""


@router.post("/applications", status_code=201)
async def create_application(
    data: ApplicationIn,
    idempotency_key: str | None = Header(default=None, alias="Idempotency-Key"),
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    job = db.scalar(select(Job).where(Job.id == data.job_id, Job.created_by == user.id))
    if not job:
        raise AppError("JOB_NOT_FOUND", "岗位不存在。", 404)
    version = owned_version(db, user, data.resume_version_id) if data.resume_version_id else None
    if version and version.job_id != job.id:
        raise AppError("RESUME_VERSION_JOB_MISMATCH", "定制简历与所选岗位不一致。", 409)
    if version and (version.guard_status != "passed" or version.review_status != "approved"):
        raise AppError(
            "RESUME_NOT_APPROVED", "定制简历通过事实检查和顾问审核后才能进入投递队列。", 409
        )
    key = idempotency_key or f"application:{user.id}:{uuid.uuid4()}"
    existing_ledger = db.scalar(
        select(QuotaLedger).where(
            QuotaLedger.user_id == user.id,
            QuotaLedger.quota_type == "application",
            QuotaLedger.idempotency_key == key,
        )
    )
    if (
        existing_ledger
        and existing_ledger.reference_type == "application"
        and existing_ledger.reference_id
    ):
        existing = db.scalar(
            select(Application).where(
                Application.id == existing_ledger.reference_id,
                Application.user_id == user.id,
            )
        )
        if existing:
            return dump(existing)
    if existing_ledger:
        raise AppError("IDEMPOTENCY_KEY_REUSED", "该幂等键已用于其他额度操作，请更换后重试。", 409)
    app = Application(
        user_id=user.id,
        job_id=job.id,
        resume_version_id=data.resume_version_id,
        source_url=job.source_url,
        notes=data.notes,
    )
    db.add(app)
    db.flush()
    change(
        db,
        user_id=user.id,
        quota_type="application",
        delta=-1,
        reason="加入投递队列",
        idempotency_key=key,
        reference_type="application",
        reference_id=app.id,
    )
    db.add(
        ApplicationEvent(
            application_id=app.id,
            from_status=None,
            to_status=app.status,
            actor_user_id=user.id,
            note="创建投递",
        )
    )
    db.commit()
    return dump(app)


def owned_application(db, user, application_id):
    row = db.scalar(
        select(Application).where(Application.id == application_id, Application.user_id == user.id)
    )
    if not row:
        raise AppError("APPLICATION_NOT_FOUND", "投递记录不存在。", 404)
    return row


@router.get("/applications")
async def applications(user: User = Depends(current_user), db: Session = Depends(get_db)):
    return listed(
        [
            dump(x)
            for x in db.scalars(
                select(Application)
                .where(Application.user_id == user.id)
                .order_by(Application.created_at.desc())
            ).all()
        ]
    )


@router.get("/applications/{application_id}")
async def application_detail(
    application_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)
):
    return dump(owned_application(db, user, application_id))


TRANSITIONS = {
    "pending_confirmation": {"preparing_resume", "withdrawn"},
    "preparing_resume": {"ready_to_apply"},
    "ready_to_apply": {"applied", "withdrawn"},
    "applied": {"contacted", "rejected"},
    "contacted": {"assessment", "interview", "rejected"},
    "assessment": {"interview", "rejected"},
    "interview": {"offer", "rejected"},
}


class StatusIn(BaseModel):
    status: str
    note: str = ""


@router.patch("/applications/{application_id}/status")
async def update_application(
    application_id: str,
    data: StatusIn,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    row = owned_application(db, user, application_id)
    if data.status not in TRANSITIONS.get(row.status, set()):
        raise AppError(
            "INVALID_STATUS_TRANSITION", f"不能从 {row.status} 变更为 {data.status}。", 409
        )
    old = row.status
    row.status = data.status
    if data.status == "applied":
        row.applied_at = datetime.now(timezone.utc)
    db.add(
        ApplicationEvent(
            application_id=row.id,
            from_status=old,
            to_status=data.status,
            actor_user_id=user.id,
            note=data.note,
        )
    )
    db.commit()
    return dump(row)


@router.post("/applications/{application_id}/confirm")
async def confirm_application(
    application_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)
):
    row = owned_application(db, user, application_id)
    if row.status != "pending_confirmation":
        raise AppError("INVALID_STATUS_TRANSITION", "该投递当前不能确认。", 409)
    row.user_confirmed_at = datetime.now(timezone.utc)
    return await update_application(
        application_id, StatusIn(status="preparing_resume", note="用户确认"), user, db
    )


@router.get("/applications/{application_id}/events")
async def application_events(
    application_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)
):
    row = owned_application(db, user, application_id)
    return listed(
        [
            dump(x)
            for x in db.scalars(
                select(ApplicationEvent)
                .where(ApplicationEvent.application_id == row.id)
                .order_by(ApplicationEvent.created_at)
            ).all()
        ]
    )


class InterviewIn(BaseModel):
    job_id: str
    resume_version_id: str | None = None


@router.post("/interviews", status_code=201)
async def create_interview(
    data: InterviewIn,
    idempotency_key: str | None = Header(default=None, alias="Idempotency-Key"),
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    job = db.scalar(select(Job).where(Job.id == data.job_id, Job.created_by == user.id))
    if not job:
        raise AppError("JOB_NOT_FOUND", "岗位不存在。", 404)
    version = owned_version(db, user, data.resume_version_id) if data.resume_version_id else None
    if version and version.job_id != job.id:
        raise AppError("RESUME_VERSION_JOB_MISMATCH", "定制简历与所选岗位不一致。", 409)
    if version and (version.guard_status != "passed" or version.review_status != "approved"):
        raise AppError("RESUME_NOT_APPROVED", "仅可使用事实检查通过且已审核的简历生成面试题。", 409)
    key = idempotency_key or f"interview:{user.id}:{uuid.uuid4()}"
    existing_ledger = db.scalar(
        select(QuotaLedger).where(
            QuotaLedger.user_id == user.id,
            QuotaLedger.quota_type == "interview",
            QuotaLedger.idempotency_key == key,
        )
    )
    if (
        existing_ledger
        and existing_ledger.reference_type == "interview"
        and existing_ledger.reference_id
    ):
        existing = db.scalar(
            select(Interview).where(
                Interview.id == existing_ledger.reference_id,
                Interview.user_id == user.id,
            )
        )
        if existing:
            return dump(existing)
    if existing_ledger:
        raise AppError("IDEMPOTENCY_KEY_REUSED", "该幂等键已用于其他额度操作，请更换后重试。", 409)
    fallback_questions = generate_questions(job, version.content_json if version else None)
    structured, provider_name = await generate_structured(
        task="interview_questions",
        system_prompt=(
            "你是模拟面试题生成助手。题目只能依据输入 JD 与已审核简历，覆盖简历真实性、"
            "岗位知识、行为问题和工作条件；不得补充未提供的公司事实。输出严格 JSON。"
        ),
        user_payload={
            "job": {
                "title": job.title,
                "responsibilities": job.responsibilities_json,
                "must_have": job.must_have_json,
                "work_conditions": job.work_conditions_json,
            },
            "approved_resume": version.content_json if version else None,
            "mock_response": fallback_questions,
        },
        response_model=InterviewStructuredOutput,
    )
    questions = structured.model_dump()
    questions["provider"] = provider_name
    row = Interview(
        user_id=user.id,
        job_id=job.id,
        resume_version_id=data.resume_version_id,
        question_set_json=questions,
    )
    db.add(row)
    db.flush()
    change(
        db,
        user_id=user.id,
        quota_type="interview",
        delta=-1,
        reason="创建模拟面试",
        idempotency_key=key,
        reference_type="interview",
        reference_id=row.id,
    )
    db.commit()
    return dump(row)


@router.get("/interviews")
async def interviews(user: User = Depends(current_user), db: Session = Depends(get_db)):
    return listed(
        [
            dump(x)
            for x in db.scalars(
                select(Interview)
                .where(Interview.user_id == user.id)
                .order_by(Interview.created_at.desc())
            ).all()
        ]
    )


@router.get("/interviews/{interview_id}")
async def interview_detail(
    interview_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)
):
    row = db.scalar(
        select(Interview).where(Interview.id == interview_id, Interview.user_id == user.id)
    )
    if not row:
        raise AppError("INTERVIEW_NOT_FOUND", "模拟面试不存在。", 404)
    return dump(row)


class AnswersIn(BaseModel):
    answers: dict


@router.patch("/interviews/{interview_id}/answers")
async def answers(
    interview_id: str,
    data: AnswersIn,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    row = db.scalar(
        select(Interview).where(Interview.id == interview_id, Interview.user_id == user.id)
    )
    if not row:
        raise AppError("INTERVIEW_NOT_FOUND", "模拟面试不存在。", 404)
    if row.status != "draft":
        raise AppError("INTERVIEW_NOT_EDITABLE", "已提交的模拟面试不能继续修改回答。", 409)
    row.answers_json = data.answers
    db.commit()
    return dump(row)


@router.post("/interviews/{interview_id}/submit")
async def submit_interview(
    interview_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)
):
    row = db.scalar(
        select(Interview).where(Interview.id == interview_id, Interview.user_id == user.id)
    )
    if not row:
        raise AppError("INTERVIEW_NOT_FOUND", "模拟面试不存在。", 404)
    if row.status != "draft":
        raise AppError("INVALID_INTERVIEW_STATUS", "该模拟面试当前不能提交。", 409)
    if not any(str(answer).strip() for answer in row.answers_json.values()):
        raise AppError("INTERVIEW_ANSWERS_REQUIRED", "请至少填写一条回答后再提交。", 409)
    row.status = "submitted"
    db.commit()
    return dump(row)


@router.get("/packages")
async def packages(db: Session = Depends(get_db)):
    return listed(
        [dump(x) for x in db.scalars(select(Package).where(Package.is_active.is_(True))).all()]
    )


@router.get("/quotas")
async def quotas(user: User = Depends(current_user), db: Session = Depends(get_db)):
    return {q: balance(db, user.id, q) for q in ["match", "resume", "application", "interview"]}


@router.get("/quotas/ledger")
async def quota_ledger(user: User = Depends(current_user), db: Session = Depends(get_db)):
    return listed(
        [
            dump(x)
            for x in db.scalars(
                select(QuotaLedger)
                .where(QuotaLedger.user_id == user.id)
                .order_by(QuotaLedger.created_at.desc())
            ).all()
        ]
    )


@router.get("/reports/weekly")
async def report(user: User = Depends(current_user), db: Session = Depends(get_db)):
    return weekly_report(db, user.id)


@router.get("/notifications")
async def notifications(user: User = Depends(current_user), db: Session = Depends(get_db)):
    rows = db.scalars(
        select(Notification)
        .where(Notification.user_id == user.id)
        .order_by(Notification.created_at.desc())
    ).all()
    return {
        **listed([dump(row) for row in rows]),
        "unread_count": sum(row.read_at is None for row in rows),
    }


@router.post("/notifications/{notification_id}/read")
async def read_notification(
    notification_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)
):
    row = db.scalar(
        select(Notification).where(
            Notification.id == notification_id, Notification.user_id == user.id
        )
    )
    if not row:
        raise AppError("NOTIFICATION_NOT_FOUND", "通知不存在。", 404)
    row.read_at = datetime.now(timezone.utc)
    db.commit()
    return dump(row)


@router.get("/account/export")
async def account_export(user: User = Depends(current_user), db: Session = Depends(get_db)):
    applications = list(db.scalars(select(Application).where(Application.user_id == user.id)))
    application_ids = [application.id for application in applications]
    payload = {
        "exported_at": datetime.now(timezone.utc).isoformat(),
        "user": dump(user, {"password_hash"}),
        "profile": dump(db.get(UserProfile, user.id)) if db.get(UserProfile, user.id) else None,
        "resumes": [
            dump(row, {"original_file_key"})
            for row in db.scalars(select(Resume).where(Resume.user_id == user.id))
        ],
        "resume_facts": [
            dump(row) for row in db.scalars(select(ResumeFact).where(ResumeFact.user_id == user.id))
        ],
        "jobs": [dump(row) for row in db.scalars(select(Job).where(Job.created_by == user.id))],
        "matches": [dump(row) for row in db.scalars(select(Match).where(Match.user_id == user.id))],
        "resume_versions": [
            dump(row, {"export_file_key"})
            for row in db.scalars(select(ResumeVersion).where(ResumeVersion.user_id == user.id))
        ],
        "applications": [dump(row) for row in applications],
        "application_events": [
            dump(row)
            for row in db.scalars(
                select(ApplicationEvent).where(ApplicationEvent.application_id.in_(application_ids))
            )
        ]
        if application_ids
        else [],
        "interviews": [
            dump(row) for row in db.scalars(select(Interview).where(Interview.user_id == user.id))
        ],
        "orders": [dump(row) for row in db.scalars(select(Order).where(Order.user_id == user.id))],
        "quota_ledger": [
            dump(row)
            for row in db.scalars(select(QuotaLedger).where(QuotaLedger.user_id == user.id))
        ],
        "notifications": [
            dump(row)
            for row in db.scalars(select(Notification).where(Notification.user_id == user.id))
        ],
        "audit_logs": [
            dump(row, {"ip_hash"})
            for row in db.scalars(select(AuditLog).where(AuditLog.actor_user_id == user.id))
        ],
    }
    return JSONResponse(
        payload,
        headers={
            "Content-Disposition": 'attachment; filename="ai-job-coach-data.json"',
            "Cache-Control": "private, no-store",
            "X-Content-Type-Options": "nosniff",
        },
    )
