import io
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from fastapi import APIRouter, Depends, File, Form, Response, UploadFile
from PIL import Image, UnidentifiedImageError
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.config import get_settings
from app.database import get_db
from app.dependencies import current_user
from app.errors import AppError
from app.models.entities import (
    AuditLog,
    Interview,
    Job,
    Resume,
    ResumeFact,
    ResumeVersion,
    User,
)
from app.providers.llm.runtime import generate_structured
from app.providers.llm.schemas import (
    InterviewStructuredOutput,
    JobStructuredOutput,
    ResumeStructuredOutput,
)
from app.providers.ocr import extract_jd_text
from app.providers.ocr.base import OCRProviderError
from app.providers.storage.local import LocalStorage
from app.services.fact_guard import guard_content
from app.services.interview_generator import generate_questions, guard_interview_guide
from app.services.job_normalizer import normalize_job
from app.services.resume_generator import generate_resume
from app.services.resume_parser import parse_resume
from app.services.resume_scorer import score_resume


router = APIRouter(prefix="/api/coach", tags=["coach"])

JD_IMAGE_TYPES = {
    "image/png": {".png"},
    "image/jpeg": {".jpg", ".jpeg"},
    "image/webp": {".webp"},
}
SOURCE_STATUSES = {"extracted", "proposed", "user_confirmed"}
MAX_IMAGE_PIXELS = 30_000_000


def _validate_jd_image(content: bytes, media_type: str, filename: str) -> None:
    suffix = Path(filename).suffix.lower()
    if media_type not in JD_IMAGE_TYPES or suffix not in JD_IMAGE_TYPES[media_type]:
        raise AppError("UNSUPPORTED_JD_IMAGE_TYPE", "JD 截图仅支持 PNG、JPG 或 WebP。", 415)
    try:
        with Image.open(io.BytesIO(content)) as image:
            width, height = image.size
            image.verify()
    except (UnidentifiedImageError, OSError, ValueError) as exc:
        raise AppError("INVALID_JD_IMAGE", "JD 截图无法读取，请重新截图后上传。", 422) from exc
    if width < 10 or height < 10 or width * height > MAX_IMAGE_PIXELS:
        raise AppError("INVALID_JD_IMAGE_SIZE", "JD 截图尺寸异常，请上传清晰的页面截图。", 422)


async def _structure_job(raw_jd: str) -> tuple[dict, str]:
    fallback = normalize_job(raw_jd)
    structured, provider_name = await generate_structured(
        task="job_normalization",
        system_prompt=(
            "请分析以下岗位 JD 文本，提取关键信息。\n"
            "要求：\n"
            "1. 识别岗位名称、公司名称、工作地点\n"
            "2. 提取岗位职责（最多 8 条）\n"
            "3. 提取任职要求（最多 8 条）\n"
            "4. 提取加分项（最多 5 条）\n"
            "5. 提取工作条件如轮班、出差等（最多 5 条）\n"
            "6. 提取关键词（技术栈、工具、证书等）\n"
            "7. 如果某些字段无法确认，放入 uncertain_fields\n"
            "8. 只提取文本中明确存在的信息，不得猜测\n"
            "输出严格 JSON。"
        ),
        user_payload={
            "raw_jd": raw_jd,
            "mock_response": {key: value for key, value in fallback.items() if key != "content_hash"},
        },
        response_model=JobStructuredOutput,
    )
    value = structured.model_dump()
    value["content_hash"] = fallback["content_hash"]
    return value, provider_name


async def _tailor_resume(facts: list[ResumeFact], job: Job) -> tuple[dict, str, list[str]]:
    fallback, _status, _claims = generate_resume(
        facts, job, allowed_statuses=SOURCE_STATUSES
    )
    structured, provider_name = await generate_structured(
        task="instant_tailored_resume",
        system_prompt=(
            "你是一位简历顾问。遵循以下启发式规则重写简历：\n"
            "1. STAR法则：每段经历改写为「行动+方法+量化结果」格式\n"
            "2. 强动词替换：将「负责」改为「主导/统筹/搭建」，将「参与」改为「独立完成/协调推动」\n"
            "3. 量化优先：保留并突出所有数字、百分比、金额\n"
            "4. 关键词三重放置：确保JD核心关键词出现在求职意向、技能栏、最近经历中\n"
            "5. ATS友好：使用标准section标题（工作经历、项目经验、技能证书、教育背景）\n"
            "6. 一岗一版：按JD相关性重新排列经历优先级\n\n"
            "严格约束：\n"
            "- 只能重写、压缩和排序输入的原简历来源行\n"
            "- 每个要点必须引用 source_fact_ids\n"
            "- 不得新增公司、岗位、日期、工具、证书、成果或数字\n"
            "- 无证据要求放入 excluded_gaps 并附 improvement_suggestions\n"
            "- 每条经历附 improvement_suggestions 标注可优化方向\n"
            "输出严格 JSON。"
        ),
        user_payload={
            "source_facts": [
                {
                    "id": fact.id,
                    "status": fact.verification_status,
                    "category": fact.category,
                    "label": fact.label,
                    "value": fact.value_json,
                }
                for fact in facts
                if fact.verification_status in SOURCE_STATUSES
            ],
            "job": {
                "title": job.title,
                "company_name": job.company_name,
                "must_have": job.must_have_json,
                "keywords": job.keywords_json,
            },
            "mock_response": fallback,
        },
        response_model=ResumeStructuredOutput,
    )
    content = structured.model_dump()
    source_facts = [fact for fact in facts if fact.verification_status in SOURCE_STATUSES]
    guard_status, claims = guard_content(content, source_facts)
    if guard_status == "blocked":
        content = fallback
        guard_status, claims = guard_content(content, source_facts)
        provider_name = f"{provider_name}_fact_guard_fallback"
    if guard_status != "passed":
        raise AppError(
            "FACT_GUARD_BLOCKED",
            "生成内容包含无法追溯的信息，已停止输出，请核对原简历后重试。",
            409,
            {"new_claims": claims},
        )
    return content, provider_name, claims


async def _create_interview_guide(
    facts: list[ResumeFact], job: Job, resume_content: dict
) -> tuple[dict, str]:
    fallback = generate_questions(
        job,
        resume_content,
        facts,
        allowed_statuses=SOURCE_STATUSES,
    )
    structured, provider_name = await generate_structured(
        task="instant_interview_guide",
        system_prompt=(
            "根据目标岗位和原简历来源生成面试练习建议。遵循以下原则：\n"
            "1. 自我介绍：控制在90秒，突出与岗位最匹配的2-3个核心经历\n"
            "2. 项目回答：使用STAR框架（情境→任务→行动→结果），每个回答引用 source_fact_ids\n"
            "3. 常见问题：回答框架要体现对岗位的理解，用已确认经历佐证\n"
            "4. 准备计划：给出具体可执行的步骤，而非泛泛建议\n"
            "严格约束：只引用 source_fact_ids，不得声称题目来自目标公司，"
            "不得补造经历、数字或结果。输出严格 JSON。"
        ),
        user_payload={
            "job": {
                "title": job.title,
                "responsibilities": job.responsibilities_json,
                "must_have": job.must_have_json,
            },
            "source_facts": [
                {"id": fact.id, "category": fact.category, "value": fact.value_json}
                for fact in facts
                if fact.verification_status in SOURCE_STATUSES
            ],
            "mock_response": fallback,
        },
        response_model=InterviewStructuredOutput,
    )
    guide = structured.model_dump()
    guard_status, _issues = guard_interview_guide(guide, facts)
    if guard_status == "blocked":
        guide = fallback
        provider_name = f"{provider_name}_fact_guard_fallback"
    return guide, provider_name


@router.post("/optimize", status_code=201)
async def optimize(
    resume_file: UploadFile = File(...),
    jd_image: UploadFile = File(...),
    jd_text: str = Form(default="", max_length=20_000),
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    settings = get_settings()
    max_bytes = settings.max_upload_mb * 1024 * 1024
    resume_content = await resume_file.read()
    image_content = await jd_image.read()
    if len(resume_content) > max_bytes or len(image_content) > max_bytes:
        raise AppError("FILE_TOO_LARGE", f"单个文件不能超过 {settings.max_upload_mb}MB。", 413)
    _validate_jd_image(
        image_content,
        jd_image.content_type or "",
        jd_image.filename or "jd-image",
    )
    parsed_resume, fact_values = parse_resume(
        resume_content,
        resume_file.content_type or "",
        resume_file.filename or "resume",
    )

    ocr_error = ""
    try:
        ocr_text, ocr_provider = extract_jd_text(image_content)
    except OCRProviderError as exc:
        ocr_text, ocr_provider, ocr_error = "", "text_fallback", str(exc)
    supplement = jd_text.strip()
    raw_jd = "\n".join(value for value in [ocr_text.strip(), supplement] if value).strip()
    if len(raw_jd) < 10:
        raise AppError(
            "JD_TEXT_NOT_RECOGNIZED",
            "没有从截图中识别出足够的 JD 文字，请上传更清晰的截图或补充粘贴 JD。",
            422,
            {"ocr_error": ocr_error} if ocr_error else {},
        )
    if not ocr_text.strip() and supplement:
        ocr_provider = "user_text_fallback"

    storage = LocalStorage()
    stored_resume_key = ""
    try:
        suffix = "." + (resume_file.filename or "resume").rsplit(".", 1)[-1].lower()
        stored_resume_key = storage.save(resume_content, suffix)
        resume = Resume(
            user_id=user.id,
            original_file_key=stored_resume_key,
            original_media_type=resume_file.content_type or "",
            display_name=(resume_file.filename or "简历")[:200],
            parse_status="parsed",
            parsed_text=parsed_resume,
        )
        db.add(resume)
        db.flush()
        facts = []
        for value in fact_values:
            fact = ResumeFact(
                resume_id=resume.id,
                user_id=user.id,
                created_by=user.id,
                source_location="用户上传的原简历",
                **value,
            )
            db.add(fact)
            facts.append(fact)
        db.flush()

        job_value, job_provider = await _structure_job(raw_jd)
        job = Job(
            created_by=user.id,
            raw_jd=raw_jd,
            source_type="user_screenshot",
            company_name=job_value["company_name"],
            title=job_value["title"],
            city=job_value["city"],
            responsibilities_json=job_value["responsibilities"],
            must_have_json=job_value["must_have"],
            nice_to_have_json=job_value["nice_to_have"],
            keywords_json=job_value["keywords"],
            work_conditions_json=job_value["work_conditions"],
            content_hash=job_value["content_hash"],
        )
        db.add(job)
        db.flush()

        tailored_content, resume_provider, new_claims = await _tailor_resume(facts, job)
        quality = score_resume(
            tailored_content,
            {
                "keywords": job.keywords_json,
                "must_have": job.must_have_json,
                "responsibilities": job.responsibilities_json,
            },
        )
        tailored_content["quality_score"] = quality["score"]
        tailored_content["quality_dimensions"] = quality["dimensions"]
        tailored_content["quality_suggestions"] = quality["suggestions"]
        version = ResumeVersion(
            user_id=user.id,
            base_resume_id=resume.id,
            job_id=job.id,
            title=f"{job.company_name} - {job.title}",
            content_json=tailored_content,
            traceability_json={
                "new_claims": new_claims,
                "provider": resume_provider,
                "source_status": "requires_user_review",
            },
            guard_status="passed",
            review_status="draft",
        )
        db.add(version)
        db.flush()

        interview_guide, interview_provider = await _create_interview_guide(
            facts, job, tailored_content
        )
        interview = Interview(
            user_id=user.id,
            job_id=job.id,
            resume_version_id=version.id,
            status="draft",
            question_set_json=interview_guide,
        )
        db.add(interview)
        db.add(
            AuditLog(
                actor_user_id=user.id,
                action="coach.optimized",
                entity_type="resume_version",
                entity_id=version.id,
                metadata_json={
                    "ocr_provider": ocr_provider,
                    "job_provider": job_provider,
                    "resume_provider": resume_provider,
                    "interview_provider": interview_provider,
                },
            )
        )
        db.commit()
        return {
            "job": {
                "id": job.id,
                "title": job.title,
                "company_name": job.company_name,
                "city": job.city,
                "responsibilities": job.responsibilities_json,
                "must_have": job.must_have_json,
                "nice_to_have": job.nice_to_have_json,
                "keywords": job.keywords_json,
                "work_conditions": job.work_conditions_json,
                "uncertain_fields": job_value["uncertain_fields"],
            },
            "tailored_resume": {
                "id": version.id,
                "title": version.title,
                "guard_status": version.guard_status,
                "content": version.content_json,
            },
            "interview_guide": interview.question_set_json,
            "ocr_provider": ocr_provider,
            "notices": [
                "JD 截图和原简历由机器提取，结果尚未经人工验证。",
                "定制简历与面试内容是建议稿，请逐项核对后使用。",
            ],
        }
    except Exception:
        db.rollback()
        if stored_resume_key:
            storage.delete(stored_resume_key)
        raise


@router.get("/resume-versions/{version_id}/download")
async def download_draft(
    version_id: str,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    version = db.scalar(
        select(ResumeVersion).where(
            ResumeVersion.id == version_id,
            ResumeVersion.user_id == user.id,
        )
    )
    if not version:
        raise AppError("RESUME_VERSION_NOT_FOUND", "定制简历不存在。", 404)
    if version.guard_status != "passed":
        raise AppError("FACT_GUARD_BLOCKED", "该建议稿未通过事实检查，不能下载。", 409)
    document = Document()
    document.add_heading(version.content_json.get("headline", version.title), 0)
    document.add_paragraph("建议稿，请根据原简历逐项核对后使用。")
    for section in version.content_json.get("sections", []):
        document.add_heading(section.get("name", "相关经历"), level=1)
        for item in section.get("items", []):
            document.add_paragraph(item.get("text", ""), style="List Bullet")
    stream = io.BytesIO()
    document.save(stream)
    db.add(
        AuditLog(
            actor_user_id=user.id,
            action="coach.resume_downloaded",
            entity_type="resume_version",
            entity_id=version.id,
            metadata_json={"format": "docx", "draft": True},
        )
    )
    db.commit()
    filename = f"tailored-resume-{datetime.now(timezone.utc):%Y%m%d}.docx"
    return Response(
        content=stream.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "Cache-Control": "private, no-store",
            "X-Content-Type-Options": "nosniff",
        },
    )
